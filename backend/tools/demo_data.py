"""GitHub Pages 시연판용 정적 샘플 데이터 생성·검증(개발·빌드 단계 전용, Pages 런타임에서는 Python을 실행하지 않음).

    cd backend && uv run python -m tools.demo_data          # 생성(frontend/demo-public/demo-data 덮어씀)
    cd backend && uv run python -m tools.demo_data --check  # 커밋된 샘플이 현재 모델·계산 엔진과 맞는지 검증

- 서버 앱을 임시 DATA_DIR로 띄워(TestClient, 샘플 규칙 모드 — AI 호출 없음) 실제 API로 두 프로젝트를 만든다.
- 화면이 쓰는 GET 응답(프로젝트·FP 집계·견적 미리보기·엑셀 미리보기·변경이력·스냅샷)과
  서버 엑셀 생성기로 만든 .xlsx를 그대로 저장한다. FP·금액을 손으로 적지 않는다.
- 가상 데이터만 쓴다(실제 사업명·개인정보·업무 문서·키 없음).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
import tempfile
from decimal import Decimal
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient

from app.config import Settings
from app.excel import export as xl
from app.fp import engine
from app.main import create_app
from app.models import Project

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "frontend" / "demo-public" / "demo-data"
SCHEMA = 1
# 엑셀·화면에 찍히는 생성 모드 표시(실제 앱의 '샘플 모드'와 구분)
DEMO_LABEL = "샘플 시연판(가상 데이터) — 서버 샘플 규칙으로 미리 생성, AI 호출 없음"
DEMO_AI = {
    "name": "demo",
    "kind": "demo",
    "label": "없음(정적 시연)",
    "model": "",
    "external": False,
    "is_real_ai": False,
    "display": "없음 — 정적 시연 데이터(AI 호출 안 함)",
    "usable": False,
    "missing": [],
    "default": True,
    "note": "이 시연판은 AI를 호출하지 않습니다. 질문·초안·FP 후보는 서버의 샘플 규칙으로 미리 만들어 둔 결과입니다.",
}

# 참고문서 샘플: 가상 기관 '가나시'의 가상 과업지시서(실제 문서 아님)
# 샘플 B 원문은 저장소 examples/의 예제 B 파일과 같은 원본(화면 다운로드·시연판·실제 앱 예제가 한 파일을 씀)
REF_DOC_PARAGRAPHS = [ln for ln in (Path(__file__).resolve().parents[2] / "examples" / "B-reference" / "reference_task_order.txt").read_text(encoding="utf-8").splitlines() if ln.strip()]


def _settings(tmp: Path) -> Settings:
    s = Settings()
    s.app_env = "development"
    s.auth_disabled = True
    s.data_dir = tmp
    s.secret_key = "demo-generator-" + "x" * 32  # 임시 디렉터리에서만 쓰는 값(배포·저장 안 함)
    s.ai_default = "sample"
    s.frontend_dist = None
    s.official_source_allowlist = ["www.postgresql.org"]
    return s


class Api:
    def __init__(self, c: TestClient):
        self.c = c

    def __call__(self, method: str, url: str, **kw) -> Any:
        r = self.c.request(method, "/api" + url, headers={"X-Requested-With": "swd"}, **kw)
        if r.status_code >= 400:
            raise RuntimeError(f"{method} {url} → {r.status_code}: {r.text[:500]}")
        return r.json() if r.headers.get("content-type", "").startswith("application/json") else r.content


def _job(api: Api, url: str, body: dict | None = None) -> dict:
    j = api("POST", url, json=body or {})["job"]
    if j["status"] != "succeeded":
        raise RuntimeError(f"{url} 작업 실패: {j['error']}")
    return j


def _apply_all(api: Api, pid: str, prid: str, skip: set[str] = frozenset()) -> None:
    pr = next(x for x in api("GET", f"/projects/{pid}")["project"]["proposals"] if x["id"] == prid)
    keys = [i["key"] for i in pr["items"] if i["status"] == "pending" and i["key"] not in skip]
    api("POST", f"/projects/{pid}/proposals/{prid}/apply", json={"keys": keys})
    if skip:
        api("POST", f"/projects/{pid}/proposals/{prid}/reject", json={"keys": sorted(skip)})


def _items(api: Api, pid: str, coll: str) -> list[dict]:
    return api("GET", f"/projects/{pid}")["project"][coll]


def _patch(api: Api, pid: str, coll: str, iid: str, data: dict) -> None:
    api("PATCH", f"/projects/{pid}/items/{coll}/{iid}", json={"data": data})


# ---------------------------------------------------------------- A. 구상 기반 — 정상 완료 사례
def build_a(api: Api) -> str:
    p = api("POST", "/projects", json={"name": "[샘플] 직원 교육 신청·관리 시스템", "sample": "education",
                                        "overview": {"budget_mode": "target", "target_amount": 45000000, "target_includes_vat": True}})["project"]
    pid = p["id"]
    # 질문(샘플 규칙) → 선택·모름·나중에 결정 예시
    _job(api, f"/projects/{pid}/ai/questions")
    qs = [q for q in _items(api, pid, "questions") if q["status"] == "open"]
    for i, q in enumerate(qs):
        if i == len(qs) - 1:
            api("POST", f"/projects/{pid}/questions/{q['id']}/answer", json={"kind": "defer"})
        elif i == 1:
            api("POST", f"/projects/{pid}/questions/{q['id']}/answer", json={"kind": "unknown"})
        else:
            api("POST", f"/projects/{pid}/questions/{q['id']}/answer", json={"kind": "choice", "value": q["recommended_value"]})
    # 답변에 따른 제안은 검토 후 거절(정상 완료 사례는 열린 변경안 없이 끝낸다)
    for pr in api("GET", f"/projects/{pid}")["project"]["proposals"]:
        if pr["status"] == "pending":
            api("POST", f"/projects/{pid}/proposals/{pr['id']}/reject", json={"keys": [i["key"] for i in pr["items"] if i["status"] == "pending"]})
    # 사용자가 추가로 적은 기능 → 요구사항 초안: 사용자 기능만 적용하고 규칙 제안은 거절(이력에 남음)
    j = _job(api, f"/projects/{pid}/ai/requirements", {"feature_text": "교육 만족도 설문: 이수 후 직원이 만족도를 입력\n교육 일정 알림: 교육 3일 전 신청자에게 안내"})
    prid = j["result"]["proposal_id"]
    pr = next(x for x in api("GET", f"/projects/{pid}")["project"]["proposals"] if x["id"] == prid)
    others = {i["key"] for i in pr["items"] if i["data"].get("origin") != "user"}
    _apply_all(api, pid, prid, skip=others)
    for r in _items(api, pid, "requirements"):
        if r["status"] != "confirmed":
            _patch(api, pid, "requirements", r["id"], {"status": "confirmed", "acceptance": r["acceptance"] or "기능 시연으로 확인"})
    # 새 요구사항의 FP — 담당자가 분류·DET·FTR을 확정한 값(엔진이 가중치·복잡도·금액 계산)
    reqs = {r["title"]: r["id"] for r in _items(api, pid, "requirements")}
    data = {d["name"]: d["id"] for d in _items(api, pid, "data_entities")}
    new_fp = [("만족도", "교육 만족도 입력", "EI", 1, 6, "직원이 입력한 만족도를 신청 정보(ILF)에 기록하는 외부 입력"),
              ("알림", "교육 일정 알림 발송", "EO", 2, 5, "대상자 선정(신청완료·교육 3일 전) 로직을 거쳐 외부(문자·메일)로 내보내는 출력")]
    for kw, name, t, ftr, det, why in new_fp:
        rid = next((v for k, v in reqs.items() if kw in k), None)
        if rid:
            api("POST", f"/projects/{pid}/items/fp_items", json={"data": {"name": name, "application": "교육관리시스템", "fp_type": t, "ftr_ret": ftr, "det": det,
                                                                            "requirement_ids": [rid], "data_ids": [data["교육 신청 정보"]], "status": "confirmed", "rationale": why}})
    # 중복 의심 해소 예시: 관리자 조회는 직원 조회와 같은 처리 → 중복으로 표시하고 제외
    fps = _items(api, pid, "fp_items")
    base = next(f for f in fps if f["name"] == "신청 내역 조회")
    dup = next(f for f in fps if f["name"].startswith("(관리자)"))
    _patch(api, pid, "fp_items", dup["id"], {"status": "excluded", "duplicate_of": base["id"], "uncertainty": "직원 조회와 처리·데이터가 같아 중복으로 판단(검토 완료)"})
    # 연계 방식 확정
    for itf in _items(api, pid, "interfaces"):
        _patch(api, pid, "interfaces", itf["id"], {"method": "인사시스템 조회 API(REST, 실시간)", "status": "confirmed"})
    # 설계 초안(샘플 규칙) → 비어 있던 화면·업무흐름만 적용(기능·데이터·연계는 이미 확정본이 있어 거절)
    j = _job(api, f"/projects/{pid}/ai/design")
    prid = j["result"]["proposal_id"]
    pr = next(x for x in api("GET", f"/projects/{pid}")["project"]["proposals"] if x["id"] == prid)
    _apply_all(api, pid, prid, skip={i["key"] for i in pr["items"] if i["collection"] not in ("screens", "workflows")})
    for coll in ("screens", "workflows", "functions"):
        for it in _items(api, pid, coll):
            if it["status"] != "confirmed":
                _patch(api, pid, coll, it["id"], {"status": "confirmed"})
    api("POST", f"/projects/{pid}/task-draft/generate")
    api("POST", f"/projects/{pid}/snapshots", json={"label": "설계 확정본"})
    # 기준 산정 → 목표예산 시나리오
    e = api("POST", f"/projects/{pid}/estimates", json={"label": "기준 산정(설계 확정)"})["estimate"]
    api("POST", f"/projects/{pid}/estimates/{e['id']}/scenarios", json={"target_amount": 45000000, "includes_vat": True, "note": "목표예산 4천5백만원(부가세 포함)"})
    # 업체 제출 문서(가상) 검토: 우리 확정 FP·견적에서 만들되 가중치 1건·부가세에 오류를 심는다
    vendor_docs(api, pid, e["id"])
    api("POST", f"/projects/{pid}/review/run", json={})
    api("PUT", f"/projects/{pid}/tech/inputs", json={"deployment": "onprem", "network": "internal", "existing_os": "Linux", "existing_db": "PostgreSQL", "ops_staff": "part_time",
                                                    "total_users": 1200, "concurrent_users": 200, "data_growth_gb_per_year": 2, "integrations": "인사시스템", "availability": "normal"})
    api("POST", f"/projects/{pid}/tech/recommend")
    return pid


def vendor_docs(api: Api, pid: str, eid: str) -> None:
    from tests.fixtures_gen import xlsx_bytes

    p = api("GET", f"/projects/{pid}")["project"]
    est = next(x for x in p["estimates"] if x["id"] == eid)
    reqs = [r for r in p["requirements"] if r["kind"] == "functional"]
    task = [f"SFR-{i + 1:03d} {r['title']} 기능을 제공하여야 한다" for i, r in enumerate(reqs)]
    task += ["ICR-001 인사시스템과 연계하여 직원 정보를 조회하여야 한다", "PER-001 동시 접속 200명에서 주요 화면 응답 3초 이내", "SFR-099 기타 필요한 기능은 협의하여 제공한다"]
    design = [f"{r['title']} 기능·화면 설계 (SFR-{i + 1:03d})" for i, r in enumerate(reqs) if i != 4]  # SFR-005 설계 누락을 심음
    design += ["인사시스템 직원정보 조회 연계 설계 (ICR-001)", "성능 목표: 동시 200명 3초 이내 (PER-001)"]
    rows = [["어플리케이션명", "세부 업무명", "단위프로세스명", "FP유형", "가중치"]]
    total = Decimal(0)
    planted = False
    for f in est["fp_snapshot"]:
        w = Decimal(str(f["simple_weight"]))
        if f["fp_type"] == "EO" and not planted:
            w, planted = Decimal("5.0"), True  # 심은 오류: EO 가중치는 5.2
        total += w
        rows.append(["교육관리시스템", "교육", f["name"], f["fp_type"], float(w)])
    rows.append(["", "", "합계", "", float(total)])
    # 업체 금액산출서: 업체 FP 합계로 현재 엔진이 계산한 금액(산식·합계는 일치), 부가세만 틀리게 심는다
    from app.fp.cost import compute_cost
    from app.models import CostSettings, Overview

    vm = compute_cost(est["standard_snapshot"], CostSettings.model_validate(est["settings_snapshot"]), Overview.model_validate(p["overview"]), total)
    f = {x["key"]: x["value"] for x in vm["factors"]}
    wrong_vat = int(vm["vat"]) + 120000  # 심은 오류: 부가세가 공급가액의 10%와 다름
    cost_rows = [["구분", "값"], ["총기능점수", float(total)], ["기능점수당 단가", vm["unit_price"]], ["규모 보정계수", float(vm["size_factor"])],
                 ["연계복잡성", float(f["integration"])], ["성능 요구수준", float(f["performance"])], ["운영환경 호환성", float(f["compatibility"])], ["보안성 요구수준", float(f["security"])],
                 ["개발원가", vm["dev_cost"]], ["이윤", vm["profit"]], ["직접경비", vm["direct_total"]], ["소프트웨어 개발비(부가세 별도)", vm["supply"]],
                 ["부가가치세", wrong_vat], ["총액", vm["supply"] + wrong_vat]]
    for name, body, ctype, kind in (
        ("가상_업체_과업지시서.txt", "\n".join(task).encode(), "text/plain", "vendor_task"),
        ("가상_업체_설계서.txt", "\n".join(design).encode(), "text/plain", "vendor_design"),
        ("가상_업체_FP산정표.xlsx", xlsx_bytes({"FP산정": rows}), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "vendor_fp"),
        ("가상_업체_금액산출서.xlsx", xlsx_bytes({"개발비": cost_rows}), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "vendor_cost"),
    ):
        api("POST", f"/projects/{pid}/documents", files={"file": (name, body, ctype)}, data={"kind": kind})


def _clear_review(api: Api, pid: str) -> None:
    """연쇄 '재검토 필요' 표시를 담당자가 확인한 상태로 정리(화면의 '검토 완료'와 같은 요청)."""
    for _ in range(3):
        p = api("GET", f"/projects/{pid}")["project"]
        left = [(c, it["id"]) for c in ("requirements", "workflows", "functions", "screens", "data_entities", "interfaces", "fp_items") for it in p[c] if it["needs_review"]]
        if not left:
            return
        for c, iid in left:
            _patch(api, pid, c, iid, {"needs_review": False, "review_reason": ""})


# ---------------------------------------------------------------- B. 참고문서 기반 — 검토 필요 예시 포함
def build_b(api: Api) -> tuple[str, bytes]:
    from tests.fixtures_gen import docx_bytes

    p = api("POST", "/projects", json={"name": "[샘플] 공공시설 예약 관리 시스템(가상 과업지시서 기반)", "overview": {
        "business_type": "system", "start_mode": "document", "budget_mode": "none",
        "idea_text": "타 기관 과업지시서를 참고해 우리 기관의 공공시설 예약 시스템을 설계한다.",
        "purpose": "체육관·회의실 예약을 온라인으로 받고 승인·이용 실적을 관리한다.",
        "users": "시민(본인인증), 시설 담당자 약 20명", "build_method": "외주 신규개발", "period_months": 5,
        "org_environment": "인터넷망 웹, 기관 결제시스템 운영 여부 확인 필요", "system_kinds": ["web", "integration"], "boundary_name": "시설예약시스템"}})["project"]
    pid = p["id"]
    doc = docx_bytes(REF_DOC_PARAGRAPHS)
    d = api("POST", f"/projects/{pid}/documents", files={"file": ("가상_타기관_과업지시서.docx", doc, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"kind": "reference"})["document"]
    _job(api, f"/projects/{pid}/documents/{d['id']}/analyze")
    for r in _items(api, pid, "reference_items"):
        t = r["text"]
        if "SFR-001" in t or "SFR-003" in t or "SFR-005" in t or "SFR-008" in t:
            dec = {"decision": "keep", "reason": "우리 기관에도 그대로 필요"}
        elif "SFR-002" in t:
            dec = {"decision": "modify", "modified_text": "SFR-002 시민은 본인인증 후 시설 이용을 예약하고, 이용일 2일 전까지 취소할 수 있어야 한다", "reason": "우리 기관 취소 규정(2일 전) 반영"}
        elif "SFR-004" in t:
            dec = {"decision": "modify", "modified_text": "SFR-004 이용료는 기관 결제시스템과 연계하여 결제하여야 한다(연계 방식 확인 필요)", "reason": "타 기관 시스템명 제거, 우리 기관 결제시스템 존재 여부 확인 필요"}
        elif "SFR-007" in t:
            dec = {"decision": "exclude", "reason": "특정 제품 지정은 우리 사업에서 하지 않음"}
        elif "사업예산" in t:
            dec = {"decision": "exclude", "reason": "타 기관 예산·기간 — 우리 사업 값은 개요에서 따로 정함"}
        elif "가나시" in t:
            dec = {"decision": "exclude", "reason": "문서 제목(타 기관명) — 요구사항 아님"}
        elif "삭제하라" in t:
            dec = {"decision": "exclude", "reason": "문서 속 지시문 — 분석 대상일 뿐 실행하지 않음"}
        elif t.strip().startswith(("1.", "2.", "3.")):
            dec = {"decision": "exclude", "reason": "목차 제목 — 요구사항 아님"}
        else:
            continue  # SFR-006 등은 결정 대기로 남겨 둔다
        api("POST", f"/projects/{pid}/reference-items/{r['id']}/decide", json=dec)
    # 담당자가 원문 문장(original_text로 보존)에 짧은 요구사항명을 붙이고 분류를 정리
    titles = {"SFR-001": ("시설 예약 일정 조회", "functional", ""), "SFR-002": ("시설 예약·취소", "functional", ""), "SFR-003": ("예약 승인·반려", "functional", ""),
              "SFR-004": ("이용료 결제 연계", "functional", ""), "SFR-005": ("이용 실적·이용률 통계", "functional", ""), "SFR-008": ("개인정보 암호화 저장", "nonfunctional", "보안")}
    for r in _items(api, pid, "requirements"):
        code = (r["original_text"] or r["title"])[:7]
        if code in titles:
            t, kind, cat = titles[code]
            pay = code == "SFR-004"
            _patch(api, pid, "requirements", r["id"], {"title": t, "kind": kind, "category": cat, "description": r["title"], "status": "review" if pay else "confirmed",
                                                       "acceptance": "결제 완료 결과가 예약 상태에 반영됨(연계 방식 확정 후 구체화)" if pay else "기능 시연으로 확인"})
    # 질문(샘플 규칙) 일부만 답변 — 미답변 질문이 남은 상태
    _job(api, f"/projects/{pid}/ai/questions")
    qs = [q for q in _items(api, pid, "questions") if q["status"] == "open"]
    if qs:
        api("POST", f"/projects/{pid}/questions/{qs[0]['id']}/answer", json={"kind": "choice", "value": qs[0]["recommended_value"]})
    # 설계·FP 후보(샘플 규칙) 적용 후 일부만 확정
    j = _job(api, f"/projects/{pid}/ai/design")
    _apply_all(api, pid, j["result"]["proposal_id"])
    j = _job(api, f"/projects/{pid}/ai/fp_candidates")
    _apply_all(api, pid, j["result"]["proposal_id"])
    return pid, doc


def finish_b(api: Api, pid: str) -> None:
    """담당자 검토: 확실한 FP만 확정(DET·FTR 입력), 나머지는 미확정·중복 의심·연계 확인 필요로 남긴다."""
    # (후보 이름, 변경 내용) — 엔진이 가중치·복잡도·금액을 계산한다
    review: dict[str, dict] = {
        "시설 예약 일정 정보": {"name": "시설 정보", "ftr_ret": 1, "det": 12, "status": "confirmed", "rationale": "담당자가 시설·운영 일정을 등록·수정하며 이 시스템이 유지하는 논리 데이터"},
        "시설 예약 정보": {"name": "예약 정보", "ftr_ret": 2, "det": 16, "status": "confirmed", "rationale": "예약·승인 이력을 이 시스템이 유지(하위 그룹 승인이력 → RET 2), 개인정보 포함"},
        "예약 정보": {"status": "excluded", "rationale": "승인·반려는 예약 정보의 상태 갱신 — 같은 논리 데이터라 별도 ILF가 아님"},
        "이용 실적 이용률 정보": {"status": "excluded", "rationale": "통계는 예약 정보에서 계산하는 값 — 별도로 유지하는 데이터가 아니므로 ILF 아님"},
        "시설 예약 일정 조회": {"ftr_ret": 2, "det": 8, "status": "confirmed", "rationale": "가공 없이 시설·예약 데이터를 조회(FTR 2)"},
        "시설 예약·취소": {"name": "시설 예약 신청", "ftr_ret": 2, "det": 10, "status": "confirmed", "rationale": "예약 정보(ILF) 등록, 시설 정보 참조(FTR 2)"},
        "예약 승인": {"ftr_ret": 1, "det": 6, "status": "confirmed", "rationale": "예약 상태 갱신"},
        "예약 반려": {"ftr_ret": 1, "det": 7, "status": "confirmed", "rationale": "예약 상태·반려 사유 갱신"},
        "이용 실적·이용률 통계": {"ftr_ret": 2, "det": 7, "status": "confirmed", "rationale": "이용률 계산 등 가공 로직을 포함한 출력"},
        "이용료 결제 연계": {"status": "draft", "uncertainty": "결제시스템 연계 방식(API 호출·화면 이동)이 정해지지 않아 EO/EQ 판단 보류 — 연계 방식 확인 후 분류"},
    }
    excluded_dup = None
    for f in _items(api, pid, "fp_items"):
        ch = review.get(f["name"])
        if ch:
            _patch(api, pid, "fp_items", f["id"], ch)
            if f["name"] == "시설 예약 정보":
                excluded_dup = f["id"]
    for f in _items(api, pid, "fp_items"):
        if f["name"] == "예약 정보" and f["status"] == "excluded" and excluded_dup:
            _patch(api, pid, "fp_items", f["id"], {"duplicate_of": excluded_dup})
    reqs = {r["title"]: r["id"] for r in _items(api, pid, "requirements")}
    data = {d["name"]: d["id"] for d in _items(api, pid, "data_entities")}
    api("POST", f"/projects/{pid}/items/fp_items", json={"data": {"name": "예약 취소", "application": "시설예약시스템", "fp_type": "EI", "ftr_ret": 1, "det": 4,
                                                                    "requirement_ids": [reqs["시설 예약·취소"]], "data_ids": [data["시설 예약 정보"]], "status": "confirmed",
                                                                    "rationale": "예약 정보의 상태를 취소로 갱신(신청과 처리 논리가 달라 별도 EI)"}})
    api("POST", f"/projects/{pid}/items/fp_items", json={"data": {"name": "결제 정보(결제시스템)", "application": "", "fp_type": "EIF", "requirement_ids": [reqs["이용료 결제 연계"]],
                                                                    "status": "draft", "rationale": "결제시스템이 유지하는 결제 결과를 참조할 가능성",
                                                                    "uncertainty": "연계 방식이 정해지면 참조 항목(DET) 입력 — 그 전에는 미확정"}})
    for coll in ("functions", "screens", "workflows"):
        for it in _items(api, pid, coll):
            _patch(api, pid, coll, it["id"], {"status": "review" if "결제" in it["name"] else "confirmed"})
    for itf in _items(api, pid, "interfaces"):
        _patch(api, pid, "interfaces", itf["id"], {"name": "결제시스템 이용료 결제 연계", "counterpart": "기관 결제시스템", "status": "review",
                                                  "method": "확인 필요(결제 API 호출 또는 결제 화면 이동)", "data_items": "예약번호, 이용료, 결제결과",
                                                  "note": "기관에 결제시스템이 있는지, 연계 표준(API 규격)이 있는지 확인 필요"})
    # 중복 의심 예시: 같은 조회가 역할 표기만 다르게 한 번 더 들어온 경우(엔진이 '중복 의심'으로 표시)
    look = next(f for f in _items(api, pid, "fp_items") if f["name"] == "시설 예약 일정 조회")
    api("POST", f"/projects/{pid}/items/fp_items", json={"data": {"name": "(관리자) 시설 예약 일정 조회", "application": "시설예약시스템", "fp_type": "EQ",
                                                                    "requirement_ids": look["requirement_ids"], "data_ids": look["data_ids"], "status": "review",
                                                                    "rationale": "관리자 화면의 일정 조회", "uncertainty": "직원용 조회와 처리·데이터가 다른지 확인 필요(같으면 중복)"}})
    _clear_review(api, pid)
    # 사용자가 추가로 적은 기능 → 요구사항 초안(검토 대기: 사용자 요구와 규칙 제안의 차이를 보여 줌)
    _job(api, f"/projects/{pid}/ai/requirements", {"feature_text": "시설 이용 후기 작성\n무단 불참(노쇼) 시 일정 기간 예약 제한"})
    api("POST", f"/projects/{pid}/estimates", json={"label": "기준 산정(일부 미확정)"})
    api("POST", f"/projects/{pid}/review/run", json={"document_ids": []})
    api("PUT", f"/projects/{pid}/tech/inputs", json={"deployment": "cloud", "network": "external", "ops_staff": "none", "total_users": 50000, "concurrent_users": 300,
                                                    "integrations": "결제시스템, 문자 발송", "availability": "high"})
    api("POST", f"/projects/{pid}/tech/recommend")


# ---------------------------------------------------------------- 내보내기
def _dump(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=False) + "\n", encoding="utf-8")


def export(app, api: Api, pid: str, key: str, prefix: str, out: Path, original: tuple[str, bytes] | None) -> dict:
    """prefix: 내려받는 파일 이름 앞부분(ASCII — 일부 브라우저가 한글 다운로드 이름을 버림, 2026-10-30 확인)."""
    ctx = app.state.ctx
    d = out / key
    resp = api("GET", f"/projects/{pid}")
    resp["ai"] = dict(DEMO_AI)
    proj = resp["project"]
    _dump(d / "project.json", resp)
    _dump(d / "fp.json", api("GET", f"/projects/{pid}/fp"))
    _dump(d / "estimate-preview.json", api("GET", f"/projects/{pid}/estimate-preview"))
    _dump(d / "changes.json", api("GET", f"/projects/{pid}/changes?limit=200"))
    snaps = api("GET", f"/projects/{pid}/snapshots")
    _dump(d / "snapshots.json", snaps)
    for s in snaps:
        _dump(d / f"diff-{s['id']}.json", api("GET", f"/projects/{pid}/snapshots/{s['id']}/diff"))
    p, _ = ctx.store.get(pid)
    std = ctx.registry.get(p.cost_settings.standard_id)
    files: dict[str, str] = {}
    for eid in [""] + [e["id"] for e in proj["estimates"]]:
        wb, meta = xl.build_workbook(p, std, ctx.store.changes(pid, 1000), DEMO_LABEL, eid or None)
        wb.properties.title = f"{p.name} — 샘플 시연판(가상 데이터)"
        wb.properties.subject = DEMO_LABEL
        tag = eid or "none"
        _dump(d / f"preview-{tag}.json", {"sheets": xl.preview(wb), "meta": {"estimate_id": meta["estimate_id"], "fp_summary": meta["fp_summary"]}})
        fname = f"{prefix}_{tag}.xlsx"
        (d / fname).write_bytes(xl.to_bytes(wb))
        files[tag] = fname
    backup = api("GET", f"/projects/{pid}/backup.json")
    (d / "backup.json").write_bytes(backup if isinstance(backup, bytes) else json.dumps(backup, ensure_ascii=False).encode())
    originals = {}
    if original:
        for doc in proj["documents"]:
            ext = doc["file_name"].rsplit(".", 1)[-1]
            fname = f"{prefix}_virtual-task-order_{doc['id']}.{ext}"
            (d / fname).write_bytes(original[1])
            originals[doc["id"]] = fname
    return {"key": key, "pid": pid, "name": proj["name"], "file_prefix": prefix, "xlsx": files, "originals": originals}


def _rename(out: Path, key: str, old: str, new: str) -> None:
    """임의 프로젝트 ID를 고정 경로 ID로 바꿔 재생성해도 시연 주소가 같게 한다."""
    for f in (out / key).glob("*.json"):
        f.write_text(f.read_text(encoding="utf-8").replace(old, new), encoding="utf-8")


def generate(out: Path) -> dict:
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    with tempfile.TemporaryDirectory() as tmp:
        app = create_app(_settings(Path(tmp)), synchronous_jobs=True)
        api = Api(TestClient(app))
        a = build_a(api)
        b, doc = build_b(api)
        finish_b(api, b)
        system = api("GET", "/system")
        system.update({"ai": dict(DEMO_AI), "mode_label": "샘플 시연판(정적 데이터)", "ai_profiles": [dict(DEMO_AI)], "default_profile": "demo",
                       "app_env": "demo", "auth_disabled": False, "official_source_allowlist": []})
        _dump(out / "system.json", system)
        samples = [export(app, api, a, "sampleA", "SAMPLE-A_education-system", out, None),
                   export(app, api, b, "sampleB", "SAMPLE-B_facility-reservation", out, ("docx", doc))]
    for s, fixed in zip(samples, ("sampleA", "sampleB")):
        _rename(out, s["key"], s["pid"], fixed)
        s["pid"] = fixed
    from app.services import examples as ex

    for name, *_ in ex.FILES:  # 예제 파일(내려받기용). 실제 앱 /api/examples와 같은 내용
        data, _ = ex.read(name)
        (out / "examples" / name).parent.mkdir(parents=True, exist_ok=True)
        (out / "examples" / name).write_bytes(data)
    manifest = {"schema": SCHEMA, "label": DEMO_LABEL, "samples": samples, "examples": ex.listing(), "files": _hashes(out)}
    _dump(out / "manifest.json", manifest)
    return manifest


def _hashes(out: Path) -> dict[str, str]:
    return {str(f.relative_to(out)): hashlib.sha256(f.read_bytes()).hexdigest()[:16] for f in sorted(out.rglob("*")) if f.is_file() and f.name != "manifest.json"}


# ---------------------------------------------------------------- 검증
def check(out: Path) -> list[str]:
    """커밋된 샘플이 현재 모델·계산 엔진·엑셀과 일치하는지 확인. 문제 목록을 돌려준다(빈 목록이면 통과)."""
    import openpyxl

    from app.fp.standards import StandardRegistry

    errs: list[str] = []
    man = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    if man.get("schema") != SCHEMA:
        errs.append("manifest schema 불일치 — 다시 생성하세요")
    for rel, h in man["files"].items():
        f = out / rel
        if not f.exists():
            errs.append(f"파일 없음: {rel}")
        elif hashlib.sha256(f.read_bytes()).hexdigest()[:16] != h:
            errs.append(f"파일이 생성 후 바뀜: {rel}")
    reg = StandardRegistry(Path(tempfile.mkdtemp()))
    for s in man["samples"]:
        d = out / s["key"]
        resp = json.loads((d / "project.json").read_text(encoding="utf-8"))
        try:
            p = Project.model_validate(resp["project"])  # 모델이 바뀌었으면 여기서 실패
        except Exception as e:  # noqa: BLE001
            errs.append(f"{s['key']}: 현재 Project 모델과 맞지 않음 — {str(e)[:300]}")
            continue
        std = reg.get(p.cost_settings.standard_id)
        fp_now = engine.summary(p, std)
        fp_saved = json.loads((d / "fp.json").read_text(encoding="utf-8"))
        for k in ("simple", "detailed"):
            now_v, saved_v = str(fp_now[k]["total_fp"]), str(fp_saved[k]["total_fp"])
            if now_v != saved_v:
                errs.append(f"{s['key']}: FP 합계({k})가 현재 엔진 계산과 다름 — 저장 {saved_v} / 현재 {now_v}")
        if str(resp["status"]["fp_total"]["simple"]) != str(fp_saved["simple"]["total_fp"]):
            errs.append(f"{s['key']}: 화면 상단 FP 합계와 FP 집계가 다름")
        ids = {it.id for c in ("requirements", "functions", "data_entities", "interfaces", "screens") for it in getattr(p, c)}
        for f in p.fp_items:
            for rid in f.requirement_ids + f.function_ids + f.data_ids:
                if rid not in ids:
                    errs.append(f"{s['key']}: {f.id}가 없는 ID {rid}를 참조")
        for e in p.estimates:
            # 견적을 당시 기준값 스냅샷으로 현재 엔진에서 다시 계산해 저장값과 비교
            from app.fp import estimates

            again = estimates.compute(p, e.standard_snapshot, e.method)
            if again["input_hash"] != e.input_hash:
                errs.append(f"{s['key']}: {e.id} 산정 후 FP·설정이 바뀜(견적 갱신 필요) — 샘플을 다시 생성하세요")
            elif int(again["money"]["total"]) != int(e.result["money"]["total"]) or str(again["fp"]["total_fp"]) != str(e.result["fp"]["total_fp"]):
                errs.append(f"{s['key']}: {e.id} FP·총액이 현재 엔진 계산과 다름 — 저장 {e.result['fp']['total_fp']}FP/{e.result['money']['total']:,}원, 현재 {again['fp']['total_fp']}FP/{int(again['money']['total']):,}원")
            wb = openpyxl.load_workbook(d / s["xlsx"][e.id], data_only=False)
            texts = " ".join(str(c.value) for row in wb["사업개요"].iter_rows(max_row=3) for c in row if c.value)
            if DEMO_LABEL not in texts or DEMO_LABEL not in str(wb.properties.subject or ""):
                errs.append(f"{s['key']}: 엑셀에 샘플 시연판 표시 없음")
            total = int(e.result["money"]["total"])
            chk = {r[0]: r for r in wb["검산"].iter_rows(values_only=True) if r and r[0]}
            if (chk.get("개발비 총액") or (None,) * 4)[3] != total:
                errs.append(f"{s['key']}: 엑셀 검산 시트에 견적 {e.id} 총액 {total:,}원(Python 확인값)이 없음")
            if not any(c.data_type == "f" for row in wb["개발비_기준"].iter_rows() for c in row):
                errs.append(f"{s['key']}: 개발비_기준 시트에 금액 수식이 없음")
            fp_cells = [c.value for row in wb["검산"].iter_rows() for c in row]
            if float(Decimal(e.result["fp"]["total_fp"])) not in [v for v in fp_cells if isinstance(v, (int, float))]:
                errs.append(f"{s['key']}: 검산 시트에 견적 FP {e.result['fp']['total_fp']}가 없음")
    return errs


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="생성하지 않고 커밋된 샘플을 검증")
    ap.add_argument("--out", type=Path, default=OUT)
    a = ap.parse_args(argv)
    if a.check:
        errs = check(a.out)
        for e in errs:
            print("✗", e)
        print("샘플 검증 통과" if not errs else f"샘플 검증 실패 {len(errs)}건 — `uv run python -m tools.demo_data`로 다시 생성")
        return 1 if errs else 0
    man = generate(a.out)
    print(f"생성 완료: {a.out} — " + ", ".join(f"{s['key']}({s['name']})" for s in man["samples"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
