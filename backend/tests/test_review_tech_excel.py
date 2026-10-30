"""검토(심은 오류 탐지), 기술 추천, 엑셀 재계산 검증."""
import io
import shutil
import subprocess

import openpyxl
import pytest

from app.excel import export as xl
from tests import fixtures_gen as fg
from tests.test_documents import up

TASK = "\n".join([
    "SFR-001 교육과정 등록 기능을 제공하여야 한다",
    "SFR-002 교육 신청 기능을 제공하여야 한다",
    "SFR-003 신청 취소 기능을 제공하여야 한다",
    "SFR-004 대기자 자동 승계 기능을 제공하여야 한다",
    "SFR-005 통계 기능 등 필요시 추가 기능을 협의하여 제공한다",
    "ICR-001 인사시스템과 연계하여야 한다",
    "PER-001 동시 접속 200명에서 응답 시간 3초 이내",
])
DESIGN = "\n".join(["FN-01 교육과정 등록 (SFR-001)", "FN-02 교육 신청 (SFR-002)", "FN-03 신청 취소 (SFR-003)", "FN-04 통계 (SFR-005)", "FN-05 연계 (ICR-001)", "PER-001 성능", "FN-09 (SFR-009) 문자 발송"])
FP_ROWS = [
    ["어플리케이션명", "세부 업무명", "단위프로세스명", "FP유형", "가중치"],
    ["교육", "과정", "교육과정 정보", "ILF", 7.5],
    ["교육", "과정", "교육과정 등록", "EI", 5],  # 심은 오류: EI는 4.0
    ["교육", "신청", "교육 신청", "EI", 4],
    ["교육", "신청", "신청 취소", "EI", 4],
    ["교육", "신청", "교육 신청", "EI", 4],  # 심은 중복
    ["교육", "통계", "통계", "EO", 5.2],
    ["", "", "합계", "", 40],  # 심은 합계 오류: 실제 29.7
]
COST_ROWS = [
    ["구분", "값"],
    ["총기능점수", 29.7],
    ["기능점수당 단가", 605784],
    ["규모 보정계수", 1.28], ["연계복잡성", 0.94], ["성능 요구수준", 0.95], ["운영환경 호환성", 1.0], ["보안성 요구수준", 1.0],
    ["개발원가", 20565330],  # 29.7×605,784=17,991,784.8 ×1.28×0.94×0.95 = 20,565,329.70 → 일치(10원 이내)
    ["이윤", 4113066],
    ["직접경비", 0],
    ["소프트웨어 개발비(부가세 별도)", 24678396],  # 20,565,330+4,113,066 → 일치
    ["부가가치세", 2600000],  # 심은 오류: 10%면 2,467,839.6
    ["총액", 27278396],  # 24,678,396+2,600,000 → 일치
    ["서버 장비 구매", 0],
]


def test_review_finds_planted_errors(client):
    c = client
    pid = c.post("/api/projects", json={"name": "검토", "overview": {"start_mode": "review"}}).json()["project"]["id"]
    c.post(f"/api/projects/{pid}/items/requirements", json={"data": {"title": "대기자 자동 승계"}})
    c.post(f"/api/projects/{pid}/items/requirements", json={"data": {"title": "이수증 출력"}})
    up(c, pid, "과업지시서.txt", TASK.encode(), "vendor_task")
    up(c, pid, "설계서.txt", DESIGN.encode(), "vendor_design")
    up(c, pid, "FP산정표.xlsx", fg.xlsx_bytes({"FP산정": FP_ROWS}), "vendor_fp")
    up(c, pid, "금액산출서.xlsx", fg.xlsx_bytes({"개발비": COST_ROWS}), "vendor_cost")
    p = c.post(f"/api/projects/{pid}/review/run", json={}).json()["project"]
    f = p["reviews"]

    def find(cat, verdict, contains=""):
        return [x for x in f if x["category"] == cat and x["verdict"] == verdict and contains in (x["original_value"] + x["location"] + x["reason"])]

    miss = find("요구사항 누락", "diff", "SFR-004")
    assert miss and "과업지시서.txt" in miss[0]["location"] and miss[0]["question_to_vendor"]
    assert find("범위 확인", "check", "SFR-009")
    w = find("FP 가중치", "diff", "교육과정 등록")
    assert w and "!E3" in w[0]["location"]
    assert find("FP 중복", "check", "교육 신청")
    tot = find("FP 합계", "diff")
    assert tot and "40" in tot[0]["original_value"] and "29.7" in tot[0]["compared_value"]
    assert find("부가세", "diff")
    assert find("합계", "match")  # 공급가 합계·총액은 일치
    assert find("금액 산식", "match")  # 개발원가 산식 일치
    assert find("단가", "match")
    assert find("장비 포함 여부", "check")
    assert find("범위 불명확", "check", "SFR-005")
    conn = find("연계 조건", "check")
    assert conn and conn[0]["doc_check_only"] is False
    perf = find("성능·보안", "check")
    assert perf and perf[0]["doc_check_only"] is False
    ours = find("요구사항 누락", "diff", "이수증 출력")
    assert ours and ours[0]["location"].startswith("프로젝트 REQ-")
    assert p["review_run"]["summary"]["diff"] >= 4


def test_review_without_docs_is_insufficient_not_error(client):
    pid = client.post("/api/projects", json={"name": "빈"}).json()["project"]["id"]
    p = client.post(f"/api/projects/{pid}/review/run", json={}).json()["project"]
    assert p["reviews"] and all(x["verdict"] == "insufficient" for x in p["reviews"])


def test_vendor_docs_do_not_ask_for_item_extraction(client):
    """업체 문서는 검토에서 바로 대조 — '분석이 끝나지 않은 문서' 안내는 참고문서에만(2026-10-30 시연 데이터 생성 중 발견)."""
    c = client
    pid = c.post("/api/projects", json={"name": "검토", "overview": {"start_mode": "review"}}).json()["project"]["id"]
    d = up(c, pid, "과업지시서.txt", TASK.encode(), "vendor_task")
    assert d["extraction"]["analysis_status"] == "not_started" and d["extraction"]["chunks_total"] > 0
    nxt = c.get(f"/api/projects/{pid}").json()["status"]["next_steps"]
    assert not any("분석이 끝나지 않은" in n["text"] for n in nxt)
    up(c, pid, "참고.txt", "SFR-001 교육 신청 기능을 제공하여야 한다\n".encode(), "reference")
    nxt = c.get(f"/api/projects/{pid}").json()["status"]["next_steps"]
    assert any("분석이 끝나지 않은" in n["text"] for n in nxt)


def test_tech_recommend_onprem_and_cloud(client):
    c = client
    pid = c.post("/api/projects", json={"name": "t", "sample": "education"}).json()["project"]["id"]
    c.put(f"/api/projects/{pid}/tech/inputs", json={"deployment": "onprem", "concurrent_users": 200, "data_growth_gb_per_year": 20, "file_storage": True, "existing_db": "Oracle 19c", "ops_staff": "part_time"})
    p = c.post(f"/api/projects/{pid}/tech/recommend").json()["project"]
    recs = p["tech"]["recommendations"]
    server = next(r for r in recs if r["area"].startswith("서버 구성"))
    assert "vCPU" in server["proposal"] and "보장하지 않음" in server["assumptions"]
    db = next(r for r in recs if r["area"] == "DB")
    assert db["paid_item"] and "개발비 미포함" in db["license"]
    assert all(r["support_status"].startswith("현재 상태 미확인") for r in recs)
    c.put(f"/api/projects/{pid}/tech/inputs", json={"deployment": "cloud", "availability": "high"})
    p = c.post(f"/api/projects/{pid}/tech/recommend").json()["project"]
    recs = p["tech"]["recommendations"]
    assert not any("vCPU" in r["proposal"] or "메모리" in r["proposal"] for r in recs)
    assert any("관리형 DB" in r["proposal"] for r in recs)
    p = c.post(f"/api/projects/{pid}/official-sources/suggested").json()["project"]
    assert all(s["status"] == "registered_unverified" and not s["checked_at"] for s in p["official_sources"])
    # 사용자가 확인 후 지원 상태 입력 → 추천에 확인일과 함께 표시
    sid = next(s["id"] for s in p["official_sources"] if "PostgreSQL" in s["title"])
    c.patch(f"/api/projects/{pid}/items/official_sources/{sid}", json={"data": {"status": "verified_by_user", "checked_at": "2026-10-29", "support_status": "17.x 지원 중(사용자 확인)"}})
    p = c.post(f"/api/projects/{pid}/tech/recommend").json()["project"]
    db = next(r for r in p["tech"]["recommendations"] if r["area"] == "DB")
    assert "확인일 2026-10-29" in db["support_status"]
    # 허용 목록 밖 URL 조회 차단
    bad = c.post(f"/api/projects/{pid}/items/official_sources", json={"data": {"title": "x", "url": "https://169.254.169.254/latest"}}).json()["item"]["id"]
    assert c.post(f"/api/projects/{pid}/official-sources/{bad}/fetch").status_code == 422


def test_excel_text_not_formula_and_not_run_sheets(client):
    c = client
    pid = c.post("/api/projects", json={"name": "엑셀", "sample": "education"}).json()["project"]["id"]
    c.post(f"/api/projects/{pid}/items/requirements", json={"data": {"title": "=HYPERLINK(\"http://x\",\"클릭\")", "description": "+1+1", "acceptance": "@SUM(A1)"}})
    r = c.get(f"/api/projects/{pid}/export.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["요구사항"]
    cell = next(c_ for row in ws.iter_rows() for c_ in row if isinstance(c_.value, str) and c_.value.startswith("=HYPERLINK"))
    assert cell.data_type == "s"
    assert any("미실행" in str(c_.value) for row in wb["개발비_기준"].iter_rows() for c_ in row)
    assert any("근거 부족" in str(c_.value) or "미분류" in str(c_.value) for row in wb["FP_정통법"].iter_rows() for c_ in row)
    assert any(str(c_.value).startswith("부분 산정") for row in wb["FP_정통법"].iter_rows() for c_ in row)
    assert wb["요구사항"].freeze_panes and wb["요구사항"].auto_filter.ref
    assert wb["요구사항"].page_setup.orientation == "landscape"


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음(재계산 검증은 로컬에서 실행)")
def test_excel_recalculated_matches_python(client, tmp_path):
    c = client
    pid = c.post("/api/projects", json={"name": "재계산", "sample": "complaint"}).json()["project"]["id"]
    c.put(f"/api/projects/{pid}/cost-settings", json={"method": "detailed", "integration_level": 2, "performance_level": 3, "rounding": {"unit": 10, "mode": "half_up"}})
    e = c.post(f"/api/projects/{pid}/estimates", json={}).json()["estimate"]
    c.post(f"/api/projects/{pid}/estimates/{e['id']}/scenarios", json={"target_amount": 80_000_000, "includes_vat": True})
    c.post(f"/api/projects/{pid}/estimates/{e['id']}/scenarios", json={"target_amount": 50_000_000, "includes_vat": False})
    src = tmp_path / "out.xlsx"
    src.write_bytes(c.get(f"/api/projects/{pid}/export.xlsx").content)
    outdir = tmp_path / "recalc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)], check=True, timeout=180, capture_output=True)
    wb = openpyxl.load_workbook(outdir / "out.xlsx", data_only=True)
    rows = [r for r in wb["검산"].iter_rows(min_row=4, values_only=True) if r[0] and r[0] != "항목"]
    assert len(rows) >= 13
    for label, sheet, fval, pval, diff, verdict in rows:
        assert verdict == "일치", (label, fval, pval)
    total_row = next(r for r in wb["개발비_기준"].iter_rows(values_only=True) if r[0] == "총액(부가세 포함)")
    assert total_row[1] == total_row[2] == e["result"]["money"]["total"]
    assert xl.SHEET_NAMES == wb.sheetnames
