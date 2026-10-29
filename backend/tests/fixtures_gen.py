"""테스트용 가상 문서 생성기(실데이터 아님)."""
from __future__ import annotations

import io
import struct
import zipfile
import zlib


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None, merge: bool = False) -> bytes:
    import docx

    d = docx.Document()
    for t in paragraphs:
        d.add_paragraph(t)
    if table:
        tbl = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                tbl.cell(r, c).text = v
        if merge:
            tbl.cell(0, 0).merge(tbl.cell(0, 1))
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


def xlsx_bytes(rows_by_sheet: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in rows_by_sheet.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def hwpx_bytes(paragraphs: list[str], table: list[list[str]] | None = None, encrypted: bool = False) -> bytes:
    ns = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section"'
    body = "".join(f"<hp:p><hp:run><hp:t>{t}</hp:t></hp:run></hp:p>" for t in paragraphs)
    if table:
        rows = "".join("<hp:tr>" + "".join(f'<hp:tc><hp:cellSpan colSpan="1" rowSpan="1"/><hp:subList><hp:p><hp:run><hp:t>{c}</hp:t></hp:run></hp:p></hp:subList></hp:tc>' for c in row) + "</hp:tr>" for row in table)
        body += f"<hp:p><hp:run><hp:tbl>{rows}</hp:tbl></hp:run></hp:p>"
    xml = f'<?xml version="1.0" encoding="UTF-8"?><hs:sec {ns}>{body}</hs:sec>'
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", xml)
        manifest = "<manifest>" + ("<encryption-data/>" if encrypted else "") + "</manifest>"
        z.writestr("META-INF/manifest.xml", manifest)
    return bio.getvalue()


def _pdf(objs: list[bytes]) -> bytes:
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    x = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offs:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode())
    return out.getvalue()


def pdf_text_bytes(lines: list[str]) -> bytes:
    """ASCII 텍스트 레이어가 있는 1쪽 PDF."""
    content = "BT /F1 11 Tf 50 750 Td 14 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
    return _pdf([
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ])


def pdf_scanned_bytes() -> bytes:
    """텍스트 없이 이미지만 있는 1쪽 PDF(스캔본 흉내)."""
    img = bytes([200] * 30)  # 10x1 RGB
    content = "q 100 0 0 10 50 700 cm /Im1 Do Q"
    return _pdf([
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /XObject << /Im1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream".encode(),
        f"<< /Type /XObject /Subtype /Image /Width 10 /Height 1 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length {len(img)} >>\nstream\n".encode() + img + b"\nendstream",
    ])


def pdf_encrypted_bytes() -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter()
    for pg in PdfReader(io.BytesIO(pdf_text_bytes(["secret SFR-001 text"]))).pages:
        w.add_page(pg)
    w.encrypt(user_password="pw-1234", owner_password="owner-1234")
    bio = io.BytesIO()
    w.write(bio)
    return bio.getvalue()


def hwp_para_record(text: str) -> bytes:
    """HWP 5.0 HWPTAG_PARA_TEXT 레코드(압축 해제 상태)."""
    payload = text.encode("utf-16le") + struct.pack("<H", 13)
    hdr = 67 | (0 << 10) | (len(payload) << 20)
    return struct.pack("<I", hdr) + payload


def deflate_raw(b: bytes) -> bytes:
    c = zlib.compressobj(9, zlib.DEFLATED, -15)
    return c.compress(b) + c.flush()
