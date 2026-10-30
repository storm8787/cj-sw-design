"""API 전체 흐름(샘플 모드, AI 키 없음). 실제 AI 검증이 아니다."""
import io
import json

import openpyxl

from app.excel.export import SHEET_NAMES
from tests.conftest import client_for

EDU_FEATURES = "교육과정 등록·수정·삭제\n교육 신청: 직원이 신청기간 내 신청\n신청 취소: 교육 시작 전까지\n신청 승인·반려\n대기자 자동 승계\n과정별 신청·이수 통계\n신청 내역 조회"


def pj(r):
    assert r.status_code == 200, r.text
    return r.json()


def test_education_full_flow(client):
    c = client
    d = pj(c.post("/api/projects", json={"name": "교육 신청·관리", "overview": {"idea_text": "교육 신청·관리 프로그램이 필요하다", "boundary_name": "교육관리시스템", "system_kinds": ["web"]}}))
    pid = d["project"]["id"]
    # 질문(소수씩) — 교육 목적에 맞는 질문 포함
    pj(c.post(f"/api/projects/{pid}/ai/questions", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    assert 1 <= len(p["questions"]) <= 5
    for _ in range(4):
        pj(c.post(f"/api/projects/{pid}/ai/questions", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    keys = {q["topic"].split(":")[0] for q in p["questions"]}
    assert {"approval", "cancel", "capacity", "completion"} <= keys
    assert "batch_schedule" not in keys and "deadline" not in keys
    qs = {q["topic"].split(":")[0]: q for q in p["questions"]}
    # 선택 답변 → 요구사항 변경안, 모름 → 가정 기록, 나중에 결정
    r = pj(c.post(f"/api/projects/{pid}/questions/{qs['capacity']['id']}/answer", json={"kind": "choice", "value": "waitlist"}))
    assert r["result"]["proposal_id"]
    r = pj(c.post(f"/api/projects/{pid}/questions/{qs['personal']['id']}/answer", json={"kind": "unknown"}))
    assert r["result"]["assumption_id"]
    pj(c.post(f"/api/projects/{pid}/questions/{qs['auth']['id']}/answer", json={"kind": "defer"}))
    assert c.post(f"/api/projects/{pid}/questions/{qs['roles']['id']}/answer", json={"kind": "choice", "value": "없는값"}).status_code == 422

    # 요구사항 초안: 사용자 기능 모두 보존, 제안은 구분
    r = pj(c.post(f"/api/projects/{pid}/ai/requirements", json={"feature_text": EDU_FEATURES}))
    prid = r["job"]["result"]["proposal_id"]
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    prop = next(x for x in p["proposals"] if x["id"] == prid)
    user_items = [i for i in prop["items"] if i["data"]["origin"] == "user"]
    assert len(user_items) == 7
    assert any(i["data"]["origin"] == "sample" for i in prop["items"])
    pj(c.post(f"/api/projects/{pid}/proposals/{prid}/apply", json={}))
    for pr in p["proposals"]:
        if pr["kind"] == "answer":
            pj(c.post(f"/api/projects/{pid}/proposals/{pr['id']}/apply", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    titles = [r["title"] for r in p["requirements"]]
    for line in EDU_FEATURES.splitlines():
        assert any(line.split(":")[0].strip()[:6] in t for t in titles), line
    for r in p["requirements"]:
        pj(c.patch(f"/api/projects/{pid}/items/requirements/{r['id']}", json={"data": {"status": "confirmed"}}))

    # 설계 초안 → 적용
    r = pj(c.post(f"/api/projects/{pid}/ai/design", json={}))
    pj(c.post(f"/api/projects/{pid}/proposals/{r['job']['result']['proposal_id']}/apply", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    fr = [x["id"] for x in p["requirements"] if x["kind"] == "functional"]
    covered = {rid for f in p["functions"] for rid in f["requirement_ids"]}
    assert set(fr) <= covered
    assert all(i["verified"] == "unverified" for i in p["interfaces"])

    # FP 후보 → 적용(모두 초안) → 분류·DET/FTR 보완 후 확정
    r = pj(c.post(f"/api/projects/{pid}/ai/fp_candidates", json={}))
    pj(c.post(f"/api/projects/{pid}/proposals/{r['job']['result']['proposal_id']}/apply", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    assert p["fp_items"] and all(x["status"] == "draft" for x in p["fp_items"])
    fp = pj(c.get(f"/api/projects/{pid}/fp"))
    assert fp["simple"]["total_fp"] == "0"  # 확정 전에는 0
    for it in p["fp_items"]:
        t = it["fp_type"] or "EI"
        data = {"fp_type": t, "status": "confirmed", "det": 8 if t in ("EI", "EO", "EQ") else 12, "ftr_ret": 1}
        pj(c.patch(f"/api/projects/{pid}/items/fp_items/{it['id']}", json={"data": data}))
    fp = pj(c.get(f"/api/projects/{pid}/fp"))
    assert float(fp["simple"]["total_fp"]) > 0 and fp["detailed"]["complete"]

    # 기준 산정(정통법) → 목표예산 → 기능 유지 확인
    pj(c.put(f"/api/projects/{pid}/cost-settings", json={"method": "detailed"}))
    e = pj(c.post(f"/api/projects/{pid}/estimates", json={"label": "1차"}))["estimate"]
    assert e["complete"] and e["result"]["money"]["total"] > 0
    sc = pj(c.post(f"/api/projects/{pid}/estimates/{e['id']}/scenarios", json={"target_amount": 120_000_000, "includes_vat": True}))["scenario"]
    assert sc["result"]["total"] == 120_000_000
    assert sorted(sc["fp_ids"]) == sorted(x["id"] for x in e["fp_snapshot"])
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    assert len(p["fp_items"]) == len(e["fp_snapshot"])  # 기능 삭제 없음

    # 영향 표시: 요구사항 변경 → 연결 FP 재검토, 견적 갱신 필요(과거 견적 보존)
    rid = fr[0]
    pj(c.patch(f"/api/projects/{pid}/items/requirements/{rid}", json={"data": {"description": "정원 초과 시 대기 접수 포함"}}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    linked = [x for x in p["fp_items"] if rid in x["requirement_ids"]] + [x for x in p["functions"] if rid in x["requirement_ids"]]
    assert linked and all(x["needs_review"] for x in linked)
    fid = p["fp_items"][0]["id"]
    st = pj(c.patch(f"/api/projects/{pid}/items/fp_items/{fid}", json={"data": {"det": 30}}))["status"]
    assert st["estimate_states"][e["id"]] == "stale"
    e2 = pj(c.post(f"/api/projects/{pid}/estimates", json={"label": "2차"}))["estimate"]
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    assert [x["id"] for x in p["estimates"]] == [e["id"], e2["id"]]
    assert p["estimates"][0]["result"]["money"]["total"] == e["result"]["money"]["total"]

    # 과업 초안 생성·사용자 편집 보존
    pj(c.post(f"/api/projects/{pid}/task-draft/generate"))
    pj(c.patch(f"/api/projects/{pid}/task-draft/purpose", json={"text": "직접 쓴 목적"}))
    pj(c.post(f"/api/projects/{pid}/task-draft/generate"))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    sec = next(s for s in p["task_draft"]["sections"] if s["key"] == "purpose")
    assert sec["text"] == "직접 쓴 목적" and sec["user_edited"]

    # 엑셀: 시트·ID·수식·Python 확인값
    r = c.get(f"/api/projects/{pid}/export.xlsx?estimate_id={e2['id']}")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == SHEET_NAMES
    req_ids = {x["id"] for x in p["requirements"]}
    sheet_ids = {row[0].value for row in wb["요구사항"].iter_rows(min_row=5) if row[0].value}
    assert req_ids <= sheet_ids
    cost_ws = wb["개발비_기준"]
    formulas = [c_.value for row in cost_ws.iter_rows() for c_ in row if c_.data_type == "f"]
    assert any("ROUNDDOWN" in f for f in formulas) and any("LN(" in f for f in formulas)
    total_row = next(row for row in cost_ws.iter_rows() if row[0].value == "총액(부가세 포함)")
    assert total_row[9].data_type == "f"  # 총액은 수식(J열), Python 확인값은 검산 시트
    chk = {row[0].value: row for row in wb["검산"].iter_rows() if row[0].value}
    assert chk["개발비 총액"][3].value == e2["result"]["money"]["total"]
    for name in ("검토의견", "HW_SW추천"):
        assert any("미실행" in str(c_.value) for row in wb[name].iter_rows() for c_ in row), name
    prev = pj(c.get(f"/api/projects/{pid}/export/preview"))
    assert [s["name"] for s in prev["sheets"]] == SHEET_NAMES

    # 백업 → 복원(새 프로젝트, 동일 내용)
    b = c.get(f"/api/projects/{pid}/backup.json")
    rp = pj(c.post("/api/projects/restore", files={"file": ("b.json", b.content, "application/json")}))["project"]
    assert rp["id"] != pid and rp["name"].endswith("(복원)")
    assert [x["id"] for x in rp["requirements"]] == [x["id"] for x in p["requirements"]]
    assert rp["estimates"][-1]["result"] == p["estimates"][-1]["result"]
    bad = json.loads(b.content)
    bad["project"]["requirements"][0]["title"] = "변조"
    assert c.post("/api/projects/restore", files={"file": ("b.json", json.dumps(bad).encode(), "application/json")}).status_code == 422


def test_other_domain_questions_and_flow(client):
    """교육 템플릿 고정이 아님: 민원·배치 연계 시스템은 다른 질문이 나온다."""
    c = client
    pid = pj(c.post("/api/projects", json={"name": "민원", "overview": {"business_type": "portal", "idea_text": "민원 접수 후 부서 배정과 처리, 야간 배치로 행정시스템에 결과 전송", "system_kinds": ["web", "batch", "integration"], "boundary_name": "민원포털"}}))["project"]["id"]
    for _ in range(5):
        pj(c.post(f"/api/projects/{pid}/ai/questions", json={}))
    keys = {q["topic"].split(":")[0] for q in pj(c.get(f"/api/projects/{pid}"))["project"]["questions"]}
    assert {"deadline", "assign", "batch_schedule", "batch_failure", "board"} <= keys
    assert "completion" not in keys and "capacity" not in keys
    r = pj(c.post(f"/api/projects/{pid}/ai/requirements", json={"feature_text": "민원 신청\t시민이 본인확인 후 신청\n부서 배정\t담당 부서에 배정\n결과 전송 배치\t매일 02시 행정시스템 전송"}))
    pj(c.post(f"/api/projects/{pid}/proposals/{r['job']['result']['proposal_id']}/apply", json={}))
    r = pj(c.post(f"/api/projects/{pid}/ai/design", json={}))
    pj(c.post(f"/api/projects/{pid}/proposals/{r['job']['result']['proposal_id']}/apply", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    assert any("전송" in i["name"] for i in p["interfaces"])
    r = pj(c.post(f"/api/projects/{pid}/ai/fp_candidates", json={}))
    p = pj(c.get(f"/api/projects/{pid}"))["project"]
    items = next(x for x in p["proposals"] if x["id"] == r["job"]["result"]["proposal_id"])["items"]
    types = {i["data"]["name"]: i["data"]["fp_type"] for i in items}
    assert any(t == "EO" for n, t in types.items() if "배치" in n or "전송" in n)


def test_sample_projects_and_delete_isolation(client):
    c = client
    a = pj(c.post("/api/projects", json={"name": "a", "sample": "education"}))
    b = pj(c.post("/api/projects", json={"name": "b", "sample": "complaint"}))
    assert a["status"]["fp_total"]["simple"] == "61.4" and a["status"]["fp_total"]["detailed"] == "60"
    pa, pb = a["project"]["id"], b["project"]["id"]
    assert c.delete(f"/api/projects/{pa}?confirm_name=틀림").status_code == 422
    pj(c.delete(f"/api/projects/{pa}?confirm_name=a"))
    assert c.get(f"/api/projects/{pa}").status_code == 404
    assert pj(c.get(f"/api/projects/{pb}"))["project"]["name"] == "b"


def test_optimistic_lock_and_invalid_links(client):
    c = client
    d = pj(c.post("/api/projects", json={"name": "x"}))
    pid, v = d["project"]["id"], d["version"]
    pj(c.patch(f"/api/projects/{pid}", json={"name": "y", "version": v}))
    assert c.patch(f"/api/projects/{pid}", json={"name": "z", "version": v}).status_code == 409
    r = c.post(f"/api/projects/{pid}/items/fp_items", json={"data": {"name": "가짜", "requirement_ids": ["REQ-999"]}})
    assert r.status_code == 422 and "REQ-999" in r.json()["detail"]


def test_persistence_across_restart_and_interrupted_jobs(app_factory):
    app1 = app_factory()
    c1 = client_for(app1)
    pid = pj(c1.post("/api/projects", json={"name": "보존", "sample": "education"}))["project"]["id"]
    app1.state.ctx.db.execute("INSERT INTO jobs(id,project_id,kind,input_hash,status,progress,created_at,updated_at) VALUES('JOB-x',?,'design','h','running',0.5,'t','t')", (pid,))
    app2 = app_factory()  # 같은 DATA_DIR로 재시작
    c2 = client_for(app2)
    assert pj(c2.get(f"/api/projects/{pid}"))["project"]["name"] == "보존"
    j = pj(c2.get("/api/jobs/JOB-x"))
    assert j["status"] == "interrupted"
    r = pj(c2.post("/api/jobs/JOB-x/retry"))
    assert r["status"] in ("succeeded", "failed")


def test_duplicate_click_reuses_running_job(app_factory):
    import threading

    gate = threading.Event()
    app = app_factory(synchronous=False)
    c = client_for(app)
    pid = pj(c.post("/api/projects", json={"name": "q", "overview": {"idea_text": "교육 신청"}}))["project"]["id"]
    ctx = app.state.ctx
    orig = ctx.jobs.handlers["questions"]

    def slow(jc, params):
        gate.wait(5)
        return orig(jc, params)

    ctx.jobs.handlers["questions"] = slow
    r1 = pj(c.post(f"/api/projects/{pid}/ai/questions", json={}))
    r2 = pj(c.post(f"/api/projects/{pid}/ai/questions", json={}))
    assert r1["started"] and r2["reused"] and r1["job"]["id"] == r2["job"]["id"]
    gate.set()


def test_snapshot_diff(client):
    c = client
    pid = pj(c.post("/api/projects", json={"name": "s", "sample": "education"}))["project"]["id"]
    snap = pj(c.post(f"/api/projects/{pid}/snapshots", json={"label": "v1"}))
    pj(c.post(f"/api/projects/{pid}/items/requirements", json={"data": {"title": "새 기능"}}))
    pj(c.patch(f"/api/projects/{pid}/items/requirements/REQ-001", json={"data": {"title": "교육과정 등록·관리(수정)"}}))
    d = pj(c.get(f"/api/projects/{pid}/snapshots/{snap['id']}/diff"))
    assert d["requirements"]["added"] == ["REQ-012"]
    assert d["requirements"]["changed"][0]["id"] == "REQ-001"
