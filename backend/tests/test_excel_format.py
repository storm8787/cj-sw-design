"""엑셀 업무용 서식·인쇄 설정(2026-10-30 개편). 값·산식 보존은 test_review_tech_excel·test_api_flow가 확인한다."""
import io

import openpyxl

from app.excel import export as xl


def _wb(c, pid):
    return openpyxl.load_workbook(io.BytesIO(c.get(f"/api/projects/{pid}/export.xlsx").content))


def test_print_setup_and_styles(client):
    c = client
    pid = c.post("/api/projects", json={"name": "서식", "sample": "education"}).json()["project"]["id"]
    e = c.post(f"/api/projects/{pid}/estimates", json={}).json()["estimate"]
    c.post(f"/api/projects/{pid}/estimates/{e['id']}/scenarios", json={"target_amount": 60_000_000, "includes_vat": True})
    wb = _wb(c, pid)
    assert wb.sheetnames == xl.SHEET_NAMES and wb.sheetnames[:3] == ["사업개요", "개발비_기준", "목표예산_시나리오"]
    for ws in wb.worksheets:
        assert ws.sheet_view.showGridLines is False, ws.title
        assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A4) and ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0, ws.title
        assert str(ws.print_area).endswith(f"${ws.max_row}"), (ws.title, ws.print_area)  # 실제 데이터까지만
        assert "&P" in ws.oddFooter.center.text
        assert ws["A1"].font.name == xl.FONT_NAME and ws["A1"].font.sz >= 14
    for name in ("사업개요", "개발비_기준", "목표예산_시나리오", "검산"):
        assert wb[name].page_setup.orientation == "portrait", name
    for name in ("FP_간이법", "FP_정통법", "요구사항", "기능설계"):
        ws = wb[name]
        assert ws.page_setup.orientation == "landscape" and ws.print_title_rows, name
        assert ws.freeze_panes and ws.auto_filter.ref, name
    # FP 상세 표: 데이터 행에는 병합이 없다(필터·정렬·행 추가를 깨지 않음)
    ws = wb["FP_간이법"]
    fr = int(ws.auto_filter.ref.split(":")[1][1:].lstrip("ABCDEFGHIJKL")) if ws.auto_filter.ref else 0
    hdr = int(ws.auto_filter.ref.split(":")[0][1:])
    for m in ws.merged_cells.ranges:
        assert not (hdr < m.min_row <= fr), f"데이터 행 병합 {m}"
    # 개발비: 총액은 수식이고 강조(합계 채움), 금액 서식은 천 단위
    cost = wb["개발비_기준"]
    total = next(r for r in cost.iter_rows() if r[0].value == "총액(부가세 포함)")
    assert total[9].data_type == "f" and total[9].number_format == xl.WON and total[9].fill.fgColor.rgb.endswith("D9D9D9")
    # 목표예산은 참고값 표시
    assert "참고" in wb["목표예산_시나리오"]["A1"].value


def test_empty_project_sheets_read_naturally(client):
    c = client
    pid = c.post("/api/projects", json={"name": "빈 사업"}).json()["project"]["id"]
    wb = _wb(c, pid)
    texts = " ".join(str(v) for row in wb["사업개요"].iter_rows(values_only=True) for v in row if v)
    assert "기준 산정 미실행" in texts and "미산정" in texts
    assert "미실행" in str(wb["개발비_기준"]["A4"].value)
    assert wb["FP_간이법"].print_area


import shutil  # noqa: E402
import subprocess  # noqa: E402

import pytest  # noqa: E402


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")
def test_detailed_counts_match_when_complexity_missing(client, tmp_path):
    """업무 확인했지만 DET·FTR 근거가 없는 EI: 정통법 유형별 기능 수(수식)도 Python처럼 복잡도 판정된 항목만 센다."""
    c = client
    pid = c.post("/api/projects", json={"name": "근거 부족"}).json()["project"]["id"]
    c.post(f"/api/projects/{pid}/items/fp_items", json={"data": {"name": "신청", "fp_type": "EI", "scope_confirmed": True}})
    src = tmp_path / "x.xlsx"
    src.write_bytes(c.get(f"/api/projects/{pid}/export.xlsx").content)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path / "r"), str(src)], check=True, timeout=180, capture_output=True)
    wb = openpyxl.load_workbook(tmp_path / "r" / "x.xlsx", data_only=True)
    rows = [r for r in wb["검산"].iter_rows(min_row=5, values_only=True) if r[0]]
    bad = [r for r in rows if r[5] != "일치"]
    assert rows and not bad, bad
