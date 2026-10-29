"""업로드 문서 검증·텍스트 추출.

원칙: 읽지 못한 내용을 읽은 척하지 않는다. 형식별 추출 품질과 한계를 messages에 남긴다.
- PDF: pypdf 텍스트 레이어. 스캔(텍스트 없음)·암호화·손상 탐지. OCR 미구현.
- DOCX: 문단·표(병합 셀 탐지), 이미지 속 글자는 읽지 않음.
- XLSX/XLSM: 셀 값(수식은 저장된 캐시값). 매크로는 실행하지 않음.
- HWPX: 압축 XML 직접 파싱(문단·표).
- HWP(5.0): 본문 레코드 텍스트만 추출하는 실험 기능. 실제 HWP 샘플 검증 전 — '부분 지원(실험)'.
- TXT/MD: UTF-8 → CP949 순서로 해독.
"""
from __future__ import annotations

import io
import re
import struct
import unicodedata
import zipfile
import zlib
from dataclasses import dataclass, field
from pathlib import PurePath
from xml.etree import ElementTree as ET

ALLOWED = {"pdf", "hwpx", "hwp", "docx", "xlsx", "xlsm", "txt", "md"}
CHUNK_CHARS = 4000


class UploadRejected(ValueError):
    pass


@dataclass
class Extracted:
    status: str = "ok"  # ok | partial | failed | unsupported | scanned | encrypted
    quality: str = "high"
    messages: list[str] = field(default_factory=list)
    pages: int | None = None
    segments: list[tuple[str, str]] = field(default_factory=list)  # (위치, 텍스트)

    @property
    def chars(self) -> int:
        return sum(len(t) for _, t in self.segments)


def safe_filename(name: str) -> str:
    name = unicodedata.normalize("NFC", PurePath(name.replace("\\", "/")).name)
    name = re.sub(r"[\x00-\x1f<>:\"|?*]", "_", name).strip(" .")
    if not name or name in (".", ".."):
        raise UploadRejected("파일명이 올바르지 않습니다.")
    return name[:180]


def check_upload(name: str, data: bytes, max_bytes: int, max_uncompressed: int) -> str:
    """형식·크기·내용 서명·압축 해제 용량을 검사하고 확장자를 반환."""
    if not data:
        raise UploadRejected("빈 파일입니다.")
    if len(data) > max_bytes:
        raise UploadRejected(f"파일이 너무 큽니다(최대 {max_bytes // (1024 * 1024)}MB).")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in ALLOWED:
        raise UploadRejected(f"지원하지 않는 형식입니다: .{ext} (지원: {', '.join(sorted(ALLOWED))})")
    if ext == "pdf" and not data.startswith(b"%PDF"):
        raise UploadRejected("PDF 형식이 아닙니다(파일 내용 서명 불일치).")
    if ext in ("docx", "xlsx", "xlsm", "hwpx"):
        if not data.startswith(b"PK"):
            raise UploadRejected("압축 기반 문서 형식이 아닙니다(손상되었거나 확장자 불일치).")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                infos = z.infolist()
                if len(infos) > 5000:
                    raise UploadRejected("압축 항목 수가 너무 많습니다.")
                total = sum(i.file_size for i in infos)
                if total > max_uncompressed:
                    raise UploadRejected(f"압축 해제 용량이 너무 큽니다(최대 {max_uncompressed // (1024 * 1024)}MB).")
                for i in infos:
                    if i.filename.startswith("/") or ".." in PurePath(i.filename).parts:
                        raise UploadRejected("압축 내부 경로가 올바르지 않습니다.")
                    if i.compress_size and i.file_size / max(i.compress_size, 1) > 200:
                        raise UploadRejected("압축률이 비정상적으로 높은 항목이 있습니다(압축 폭탄 의심).")
        except zipfile.BadZipFile as e:
            raise UploadRejected("압축 구조가 손상된 파일입니다.") from e
    if ext == "hwp" and not data.startswith(b"\xd0\xcf\x11\xe0"):
        raise UploadRejected("HWP(5.0) 형식이 아닙니다. HWP 3.x 등 구형 파일은 HWPX/PDF로 변환해 주세요.")
    return ext


def extract(ext: str, data: bytes) -> Extracted:
    try:
        return {
            "pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "xlsm": _xlsx, "hwpx": _hwpx, "hwp": _hwp, "txt": _text, "md": _text,
        }[ext](data)
    except Exception as e:  # noqa: BLE001 - 손상 파일은 실패로 기록
        return Extracted(status="failed", quality="none", messages=[f"파일을 읽지 못했습니다(손상 또는 형식 불일치): {type(e).__name__}"])


def chunk(segments: list[tuple[str, str]], size: int = CHUNK_CHARS) -> list[tuple[str, str]]:
    """세그먼트를 위치 범위가 붙은 구간으로 묶는다. 각 줄 앞에 [위치]를 남겨 원문 위치를 보존."""
    out: list[tuple[str, str]] = []
    buf: list[str] = []
    first = last = ""
    n = 0
    for loc, text in segments:
        for piece in [text[i:i + size] for i in range(0, max(len(text), 1), size)]:
            line = f"[{loc}] {piece}"
            if n + len(line) > size and buf:
                out.append((first if first == last else f"{first} ~ {last}", "\n".join(buf)))
                buf, n = [], 0
            if not buf:
                first = loc
            buf.append(line)
            last = loc
            n += len(line) + 1
    if buf:
        out.append((first if first == last else f"{first} ~ {last}", "\n".join(buf)))
    return out


# ---------------------------------------------------------------- 형식별
def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        r = PdfReader(io.BytesIO(data))
    except PdfReadError as e:
        return Extracted(status="failed", quality="none", messages=[f"PDF 구조 손상: {str(e)[:120]}"])
    if r.is_encrypted:
        try:
            ok = r.decrypt("")
        except Exception:  # noqa: BLE001
            ok = 0
        if not ok:
            return Extracted(status="encrypted", quality="none", messages=["암호화된 PDF입니다. 암호를 해제한 파일을 올려 주세요."])
    ex = Extracted(pages=len(r.pages))
    empty_pages = []
    image_pages = []
    for i, page in enumerate(r.pages, 1):
        try:
            t = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001
            t = ""
            ex.messages.append(f"{i}쪽 텍스트 추출 오류")
        has_img = False
        try:
            res = page.get("/Resources") or {}
            xo = res.get("/XObject") if hasattr(res, "get") else None
            if xo:
                xo = xo.get_object()
                has_img = any(xo[k].get_object().get("/Subtype") == "/Image" for k in xo)
        except Exception:  # noqa: BLE001
            pass
        if len(t) < 20:
            empty_pages.append(i)
            if has_img:
                image_pages.append(i)
        if t:
            for j, para in enumerate([x for x in re.split(r"\n\s*\n", t) if x.strip()], 1):
                ex.segments.append((f"{i}쪽 문단{j}", " ".join(para.split())))
    if not ex.segments:
        ex.status, ex.quality = "scanned", "none"
        ex.messages.append("텍스트 레이어가 없습니다(스캔 PDF로 보임). OCR은 구현되지 않아 내용을 읽지 않았습니다. 원본 HWPX/DOCX 또는 텍스트 PDF를 올려 주세요.")
    elif empty_pages:
        ex.status, ex.quality = "partial", "medium"
        ex.messages.append(f"텍스트가 없거나 매우 적은 쪽: {_ranges(empty_pages)}" + (f" (이미지 포함 쪽 {_ranges(image_pages)} — 이미지 속 글자·표는 읽지 않음)" if image_pages else ""))
    else:
        ex.quality = "medium"
    if ex.segments:
        ex.messages.append("PDF 표는 행·열 구조가 보존되지 않고 텍스트 순서로 추출됩니다.")
    return ex


def _docx(data: bytes) -> Extracted:
    import docx

    d = docx.Document(io.BytesIO(data))
    ex = Extracted()
    pi = 0
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            pi += 1
            ex.segments.append((f"문단{pi}", t))
    merged_tables = []
    for ti, tbl in enumerate(d.tables, 1):
        for ri, row in enumerate(tbl.rows, 1):
            seen = []
            cells = []
            for c in row.cells:
                if any(c._tc is s for s in seen):
                    if ti not in merged_tables:
                        merged_tables.append(ti)
                    continue
                seen.append(c._tc)
                cells.append(" ".join(c.text.split()))
            line = " | ".join(cells)
            if line.strip(" |"):
                ex.segments.append((f"표{ti} {ri}행", line))
    imgs = len(d.inline_shapes)
    if merged_tables:
        ex.messages.append(f"병합 셀이 있는 표: {', '.join(f'표{x}' for x in merged_tables)} — 병합 셀은 한 번만 기록됨(열 위치 확인 필요)")
    if imgs:
        ex.messages.append(f"이미지 {imgs}개 — 이미지 속 글자·표는 읽지 않음")
        ex.quality = "medium"
    if not ex.segments:
        ex.status, ex.quality = "partial", "low"
        ex.messages.append("추출된 텍스트가 없습니다.")
    return ex


def _xlsx(data: bytes) -> Extracted:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_vba=False)
    wbf = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=False)
    ex = Extracted()
    no_cache = 0
    merged = []
    for ws in wb.worksheets:
        wf = wbf[ws.title]
        if ws.merged_cells.ranges:
            merged.append(ws.title)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20000)):
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    fv = wf[c.coordinate].value
                    if isinstance(fv, str) and fv.startswith("="):
                        no_cache += 1
                    continue
                vals.append(f"{c.coordinate}={v}" if not isinstance(v, str) else f"{c.coordinate}={' '.join(v.split())}")
            if vals:
                ex.segments.append((f"시트 '{ws.title}' {row[0].row}행", " | ".join(vals)))
        if ws.max_row > 20000:
            ex.status = "partial"
            ex.messages.append(f"시트 '{ws.title}'는 20,000행까지만 읽음(전체 {ws.max_row}행)")
    if merged:
        ex.messages.append(f"병합 셀이 있는 시트: {', '.join(merged)} — 병합 영역은 첫 셀에만 값이 있음")
    if no_cache:
        ex.messages.append(f"계산 결과(캐시값)가 저장되지 않은 수식 셀 {no_cache}개 — Excel에서 열어 저장한 파일을 올리면 값을 읽을 수 있음")
        ex.quality = "medium"
    if b"vbaProject" in data:
        ex.messages.append("매크로가 포함된 파일입니다. 매크로는 실행하지 않았습니다.")
    if not ex.segments:
        ex.status, ex.quality = "partial", "low"
        ex.messages.append("값이 있는 셀이 없습니다.")
    return ex


HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"


def _hwpx(data: bytes) -> Extracted:
    ex = Extracted()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        if "META-INF/manifest.xml" in names and b"encryption-data" in z.read("META-INF/manifest.xml"):
            return Extracted(status="encrypted", quality="none", messages=["암호가 설정된 HWPX입니다. 암호를 해제해 다시 올려 주세요."])
        sections = sorted([n for n in names if re.match(r"Contents/section\d+\.xml$", n)], key=lambda x: int(re.findall(r"\d+", x)[-1]))
        if not sections:
            return Extracted(status="failed", quality="none", messages=["HWPX 본문(section)을 찾지 못했습니다."])
        if any(n.startswith("BinData/") for n in names):
            ex.messages.append("그림 등 바이너리 자료가 포함됨 — 그림 속 글자·표는 읽지 않음")
        merged = False
        for si, sec in enumerate(sections, 1):
            root = ET.fromstring(z.read(sec))
            tbl_nodes = set()
            for ti, tbl in enumerate(root.iter(f"{HP}tbl"), 1):
                for el in tbl.iter():
                    tbl_nodes.add(id(el))
                for ri, tr in enumerate(tbl.iter(f"{HP}tr"), 1):
                    cells = []
                    for tc in tr.iter(f"{HP}tc"):
                        span = tc.find(f"{HP}cellSpan")
                        if span is not None and (span.get("colSpan", "1") != "1" or span.get("rowSpan", "1") != "1"):
                            merged = True
                        cells.append(" ".join("".join(t.text or "" for t in tc.iter(f"{HP}t")).split()))
                    line = " | ".join(cells)
                    if line.strip(" |"):
                        ex.segments.append((f"구역{si} 표{ti} {ri}행", line))
            pi = 0
            for p in root.iter(f"{HP}p"):
                if id(p) in tbl_nodes:
                    continue
                t = " ".join("".join(x.text or "" for x in p.iter(f"{HP}t")).split())
                if t:
                    pi += 1
                    ex.segments.append((f"구역{si} 문단{pi}", t))
        if merged:
            ex.messages.append("병합 셀이 있는 표가 있음 — 열 위치 확인 필요")
    if not ex.segments:
        ex.status, ex.quality = "partial", "low"
        ex.messages.append("추출된 텍스트가 없습니다.")
    return ex


# HWP 5.0 본문 레코드 파서(실험)
HWPTAG_PARA_TEXT = 67
_EXT_CTRL = {1, 2, 3, 11, 12, 14, 15, 16, 17, 18, 21, 22, 23}
_INLINE_CTRL = {4, 5, 6, 7, 8, 9, 19, 20}


def hwp_para_texts(body: bytes) -> list[str]:
    """압축 해제된 BodyText 섹션 바이트에서 문단 텍스트 목록을 꺼낸다."""
    out = []
    pos = 0
    while pos + 4 <= len(body):
        (hdr,) = struct.unpack_from("<I", body, pos)
        pos += 4
        tag = hdr & 0x3FF
        size = (hdr >> 20) & 0xFFF
        if size == 0xFFF:
            (size,) = struct.unpack_from("<I", body, pos)
            pos += 4
        rec = body[pos:pos + size]
        pos += size
        if tag != HWPTAG_PARA_TEXT:
            continue
        chars = struct.unpack(f"<{len(rec) // 2}H", rec[: len(rec) // 2 * 2])
        buf = []
        i = 0
        while i < len(chars):
            c = chars[i]
            if c < 32:
                if c in _EXT_CTRL or c in _INLINE_CTRL:
                    if c == 9:
                        buf.append("\t")
                    i += 8
                    continue
                if c in (10, 13):
                    buf.append("\n")
                i += 1
                continue
            buf.append(chr(c))
            i += 1
        t = "".join(buf).strip()
        if t:
            out.append(t)
    return out


def _hwp(data: bytes) -> Extracted:
    import olefile

    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        if not ole.exists("FileHeader"):
            return Extracted(status="failed", quality="none", messages=["HWP 파일 헤더가 없습니다."])
        hdr = ole.openstream("FileHeader").read()
        if not hdr.startswith(b"HWP Document File"):
            return Extracted(status="failed", quality="none", messages=["HWP 5.0 문서가 아닙니다."])
        (flags,) = struct.unpack_from("<I", hdr, 36)
        compressed, encrypted, distribution = bool(flags & 1), bool(flags & 2), bool(flags & 4)
        if encrypted:
            return Extracted(status="encrypted", quality="none", messages=["암호가 설정된 HWP입니다. 암호 해제 후 HWPX 또는 PDF로 변환해 올려 주세요."])
        if distribution:
            return Extracted(status="unsupported", quality="none", messages=["배포용(보호) HWP 문서는 읽을 수 없습니다. HWPX/PDF로 변환해 주세요."])
        ex = Extracted(status="partial", quality="low", messages=["HWP 텍스트 추출은 실험 기능입니다: 표·서식 구조가 보존되지 않고 실제 HWP 샘플로 검증되지 않았습니다. 정확한 분석은 HWPX(한글에서 '다른 이름으로 저장') 또는 PDF로 변환해 올려 주세요."])
        secs = sorted([e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText" and e[1].startswith("Section")], key=lambda e: int(e[1][7:] or 0))
        for si, e in enumerate(secs, 1):
            raw = ole.openstream(e).read()
            body = zlib.decompress(raw, -15) if compressed else raw
            for pi, t in enumerate(hwp_para_texts(body), 1):
                ex.segments.append((f"구역{si} 문단{pi}", " ".join(t.split())))
        if not ex.segments:
            ex.status, ex.quality = "failed", "none"
            ex.messages.append("본문 텍스트를 찾지 못했습니다.")
        return ex
    finally:
        ole.close()


def _text(data: bytes) -> Extracted:
    ex = Extracted()
    for enc in ("utf-8-sig", "cp949"):
        try:
            s = data.decode(enc)
            if enc == "cp949":
                ex.messages.append("CP949(EUC-KR)로 해독했습니다.")
            break
        except UnicodeDecodeError:
            continue
    else:
        s = data.decode("utf-8", errors="replace")
        ex.status, ex.quality = "partial", "low"
        ex.messages.append("문자 인코딩을 알 수 없어 일부 글자가 깨졌을 수 있습니다.")
    for i, ln in enumerate(s.splitlines(), 1):
        if ln.strip():
            ex.segments.append((f"{i}줄", ln.strip()))
    if not ex.segments:
        ex.status, ex.quality = "partial", "low"
        ex.messages.append("내용이 없습니다.")
    return ex


def _ranges(nums: list[int]) -> str:
    out, s, prev = [], None, None
    for n in nums:
        if s is None:
            s = prev = n
        elif n == prev + 1:
            prev = n
        else:
            out.append(f"{s}" if s == prev else f"{s}-{prev}")
            s = prev = n
    if s is not None:
        out.append(f"{s}" if s == prev else f"{s}-{prev}")
    return ", ".join(out) + "쪽"
