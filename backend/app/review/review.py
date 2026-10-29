"""기존(업체) 문서 대조 검토 — 규칙 기반.

판정: match(확인 일치) / diff(차이 발견) / check(확인 필요) / insufficient(근거 부족).
문서가 없거나 정보가 부족한 것은 오류로 단정하지 않고 '근거 부족'으로 둔다.
성능·연계·보안의 실제 충족은 문서로 확정할 수 없음을 표시한다(doc_check_only=False).
"""
from __future__ import annotations

import io
import re
from decimal import Decimal, InvalidOperation
from typing import Any

from ..ai.sample import REQ_ID, base_noun
from ..fp.engine import complexity_of, normalize_name
from ..models import DocumentMeta, Project

KIND_LABEL = {"vendor_task": "과업지시서", "vendor_design": "설계서", "vendor_fp": "FP 산정표", "vendor_cost": "금액 산출서", "reference": "참고문서", "other": "기타 문서"}
VAGUE = re.compile(r"(\s등\s|\s등\.|기타\s|필요\s*시|협의\s*하여|추후|적절히|원활히|최적화|충분한)")
IFACE = re.compile(r"(연계|인터페이스|API|송수신)")
PERF_SEC = re.compile(r"(응답\s*시간|동시\s*접속|TPS|처리\s*성능|보안\s*인증|취약점|암호화|가용성\s*\d)")
EQUIP = re.compile(r"(장비|서버\s*구매|하드웨어|HW\s*구매|스토리지\s*구매)")

COST_LABELS = [
    ("fp", re.compile(r"(총\s*기능\s*점수|기능\s*점수\s*합계|총\s*FP)")),
    ("unit", re.compile(r"(기능\s*점수\s*당\s*단가|FP\s*당\s*단가|FP\s*단가)")),
    ("size", re.compile(r"(규모\s*보정|SW\s*규모)")),
    ("integration", re.compile(r"연계\s*복잡")),
    ("performance", re.compile(r"성능\s*(요구|보정)")),
    ("compatibility", re.compile(r"(운영\s*환경|다중\s*사이트)")),
    ("security", re.compile(r"보안\s*(성|요구)")),
    ("dev", re.compile(r"(개발\s*원가|보정\s*후\s*개발원가)")),
    ("profit", re.compile(r"^\s*이윤")),
    ("direct", re.compile(r"직접\s*경비")),
    ("supply", re.compile(r"(개발비\s*\(?\s*부가세\s*별도|공급\s*가액|부가세\s*별도)")),
    ("vat", re.compile(r"(부가\s*가치세|^\s*부가세\s*$|VAT)")),
    ("total", re.compile(r"(총\s*액|총\s*사업비|부가세\s*포함|합\s*계\s*금액)")),
]


def _num(v: Any) -> Decimal | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).replace(",", "").replace("원", "").strip()
    try:
        return Decimal(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None
    except InvalidOperation:
        return None


def _f(**kw) -> dict[str, Any]:
    base = {"category": "", "verdict": "check", "location": "", "original_value": "", "compared_value": "", "reason": "", "suggestion": "", "question_to_vendor": "", "document_ids": [], "doc_check_only": True, "origin": "system"}
    base.update(kw)
    return base


def _lines(chunks: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """구간 텍스트를 (위치, 줄) 목록으로. 각 줄의 [위치] 접두를 복원한다."""
    out = []
    for loc, text in chunks:
        for ln in text.split("\n"):
            m = re.match(r"^\[(.+?)\]\s?(.*)$", ln)
            out.append((m.group(1), m.group(2)) if m else (loc, ln))
    return out


# ---------------------------------------------------------------- FP 표 재계산(엑셀)
def parse_fp_table(data: bytes, std: dict[str, Any]) -> dict[str, Any] | None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_vba=False)
    for ws in wb.worksheets:
        hdr_row = None
        cols: dict[str, int] = {}
        for r in range(1, min(ws.max_row, 30) + 1):
            for c in range(1, min(ws.max_column, 30) + 1):
                v = ws.cell(r, c).value
                if not isinstance(v, str):
                    continue
                s = re.sub(r"[\s①-⑩]", "", v)
                if "FP유형" in s or s in ("유형", "기능유형"):
                    cols["type"], hdr_row = c, r
                elif "단위프로세스" in s or s in ("기능명", "단위프로세스명"):
                    cols.setdefault("name", c)
                elif s.startswith("DET"):
                    cols["det"] = c
                elif "RET/FTR" in s or "FTR/RET" in s or s in ("FTR", "RET"):
                    cols["ftr"] = c
                elif "가중치" in s or s in ("점수", "FP"):
                    cols.setdefault("weight", c)
            if hdr_row:
                break
        if not hdr_row:
            continue
        rows = []
        total_decl = None
        total_loc = ""
        for r in range(hdr_row + 1, ws.max_row + 1):
            t = ws.cell(r, cols["type"]).value
            if isinstance(t, str) and t.strip().upper() in ("EI", "EO", "EQ", "ILF", "EIF"):
                rows.append({
                    "row": r, "type": t.strip().upper(),
                    "name": str(ws.cell(r, cols["name"]).value or "") if "name" in cols else "",
                    "det": _num(ws.cell(r, cols["det"]).value) if "det" in cols else None,
                    "ftr": _num(ws.cell(r, cols["ftr"]).value) if "ftr" in cols else None,
                    "weight": _num(ws.cell(r, cols["weight"]).value) if "weight" in cols else None,
                    "weight_cell": f"{ws.cell(r, cols['weight']).coordinate}" if "weight" in cols else "",
                })
            else:
                for c in range(1, min(ws.max_column, 30) + 1):
                    v = ws.cell(r, c).value
                    if isinstance(v, str) and re.search(r"(합\s*계|총\s*기능\s*점수|^\s*계\s*$)", v):
                        for c2 in range(c + 1, min(ws.max_column, 40) + 1):
                            n = _num(ws.cell(r, c2).value)
                            if n is not None:
                                total_decl, total_loc = n, f"'{ws.title}'!{ws.cell(r, c2).coordinate}"
                                break
        if rows:
            method = "detailed" if "det" in cols and "ftr" in cols and any(x["det"] is not None for x in rows) else "simple"
            sw = {k: Decimal(v) for k, v in std["simple_weights"]["values"].items()}
            for x in rows:
                if method == "detailed" and x["det"] is not None and x["ftr"] is not None:
                    cx, _ = complexity_of(std, x["type"], int(x["ftr"]), int(x["det"]))
                    x["expected"] = Decimal(std["detailed_weights"]["values"][x["type"]][cx]) if cx else None
                else:
                    x["expected"] = sw[x["type"]] if method == "simple" else None
            return {"sheet": ws.title, "method": method, "rows": rows, "total_declared": total_decl, "total_loc": total_loc}
    return None


def parse_cost_values(data: bytes | None, lines: list[tuple[str, str]], is_xlsx: bool) -> dict[str, tuple[Decimal, str]]:
    found: dict[str, tuple[Decimal, str]] = {}
    if is_xlsx and data:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_vba=False)
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300)):
                for i, c in enumerate(row):
                    if not isinstance(c.value, str):
                        continue
                    for key, rx in COST_LABELS:
                        if key in found or not rx.search(c.value):
                            continue
                        for c2 in row[i + 1:]:
                            n = _num(c2.value)
                            if n is not None:
                                found[key] = (n, f"'{ws.title}'!{c2.coordinate}")
                                break
                        break
    else:
        for loc, ln in lines:
            for key, rx in COST_LABELS:
                if key in found:
                    continue
                m = rx.search(ln)
                if m:
                    nm = re.search(r"(\d[\d,]*(\.\d+)?)\s*(원|FP)?", ln[m.end():])
                    if nm:
                        n = _num(nm.group(1))
                        if n is not None:
                            found[key] = (n, loc)
    return found


def _won(d: Decimal) -> str:
    return f"{d:,.2f}".rstrip("0").rstrip(".") if d != d.to_integral() else f"{int(d):,}"


# ---------------------------------------------------------------- 메인
def run_review(p: Project, docs: list[DocumentMeta], chunks: dict[str, list[tuple[str, str]]], files: dict[str, bytes], std: dict[str, Any]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    by_kind: dict[str, list[DocumentMeta]] = {}
    for d in docs:
        by_kind.setdefault(d.kind, []).append(d)
    lines = {d.id: _lines(chunks.get(d.id, [])) for d in docs}
    for d in docs:
        if d.extraction.status in ("scanned", "encrypted", "failed", "unsupported"):
            out.append(_f(category="문서 판독", verdict="insufficient", location=d.file_name, reason=f"{KIND_LABEL[d.kind]} 텍스트를 읽지 못함({d.extraction.status}): " + "; ".join(d.extraction.messages[:2]), suggestion="텍스트가 있는 파일(HWPX/DOCX/텍스트 PDF/XLSX)로 다시 받아 검토", document_ids=[d.id]))
        elif d.extraction.status == "partial":
            out.append(_f(category="문서 판독", verdict="check", location=d.file_name, reason="일부만 읽음: " + "; ".join(d.extraction.messages[:2]), suggestion="읽지 못한 부분은 원문으로 직접 확인", document_ids=[d.id]))

    # A. 요구사항 ID 대조
    task_docs = by_kind.get("vendor_task", [])
    other_docs = by_kind.get("vendor_design", []) + by_kind.get("vendor_fp", [])
    if not task_docs:
        out.append(_f(category="요구사항 누락", verdict="insufficient", reason="과업지시서(또는 제안요청서)가 등록되지 않아 요구사항 누락 여부를 대조할 수 없습니다(오류 아님).", suggestion="문서 종류를 '업체 과업지시서'로 등록"))
    else:
        task_ids: dict[str, tuple[str, str]] = {}
        for d in task_docs:
            for loc, ln in lines[d.id]:
                for rid in REQ_ID.findall(ln):
                    task_ids.setdefault(rid, (f"{d.file_name} {loc}", d.id))
        other_ids: dict[str, tuple[str, str]] = {}
        for d in other_docs:
            for loc, ln in lines[d.id]:
                for rid in REQ_ID.findall(ln):
                    other_ids.setdefault(rid, (f"{d.file_name} {loc}", d.id))
        if not task_ids:
            out.append(_f(category="요구사항 누락", verdict="insufficient", reason="과업지시서에서 요구사항 ID(예: SFR-001)를 찾지 못해 ID 대조를 하지 못했습니다.", document_ids=[d.id for d in task_docs]))
        elif not other_docs:
            out.append(_f(category="요구사항 누락", verdict="insufficient", reason="설계서·FP 산정표가 없어 요구사항 반영 여부를 대조할 수 없습니다.", document_ids=[d.id for d in task_docs]))
        else:
            missing = [rid for rid in task_ids if rid not in other_ids]
            for rid in missing:
                loc, did = task_ids[rid]
                out.append(_f(category="요구사항 누락", verdict="diff", location=loc, original_value=rid, compared_value="설계서·FP 산정표에 없음", reason=f"과업지시서의 {rid}가 설계서·FP 산정표 어디에도 나오지 않습니다.", suggestion=f"{rid}에 해당하는 기능설계·FP 항목 추가 또는 제외 사유 명시", question_to_vendor=f"{rid} 요구사항은 어느 기능·FP 항목에 반영되었습니까?", document_ids=[did] + [d.id for d in other_docs]))
            extra = [rid for rid in other_ids if rid not in task_ids]
            for rid in extra:
                loc, did = other_ids[rid]
                out.append(_f(category="범위 확인", verdict="check", location=loc, original_value=rid, compared_value="과업지시서에 없음", reason=f"{rid}는 과업지시서에 없는 ID입니다(범위 추가 또는 오기 가능).", question_to_vendor=f"{rid}의 근거 요구사항은 무엇입니까?", document_ids=[did]))
            if not missing:
                out.append(_f(category="요구사항 누락", verdict="match", reason=f"과업지시서 요구사항 ID {len(task_ids)}건이 모두 설계서·FP 산정표에 나타납니다(ID 기준 대조).", document_ids=[d.id for d in task_docs + other_docs]))

    # B. FP 산정표 재계산
    fp_total_from_table: Decimal | None = None
    fp_docs = by_kind.get("vendor_fp", [])
    if not fp_docs:
        out.append(_f(category="FP 산정", verdict="insufficient", reason="FP 산정표가 등록되지 않아 FP 재계산을 하지 못했습니다(오류 아님)."))
    for d in fp_docs:
        if d.format not in ("xlsx", "xlsm") or d.id not in files:
            out.append(_f(category="FP 산정", verdict="insufficient", location=d.file_name, reason="FP 산정표가 엑셀이 아니어서 행 단위 재계산을 하지 못했습니다.", suggestion="엑셀 원본 제출 요청", question_to_vendor="FP 산정표 엑셀 원본을 제출해 주십시오.", document_ids=[d.id]))
            continue
        t = parse_fp_table(files[d.id], std)
        if not t:
            out.append(_f(category="FP 산정", verdict="insufficient", location=d.file_name, reason="FP 유형 열(EI/EO/EQ/ILF/EIF)을 찾지 못했습니다.", document_ids=[d.id]))
            continue
        decl_sum = sum((x["weight"] for x in t["rows"] if x["weight"] is not None), Decimal(0))
        exp_sum = sum((x["expected"] for x in t["rows"] if x["expected"] is not None), Decimal(0))
        fp_total_from_table = decl_sum
        wrong = [x for x in t["rows"] if x["weight"] is not None and x["expected"] is not None and x["weight"] != x["expected"]]
        for x in wrong[:30]:
            out.append(_f(category="FP 가중치", verdict="diff", location=f"{d.file_name} '{t['sheet']}'!{x['weight_cell']}", original_value=f"{x['name']} {x['type']} 가중치 {_won(x['weight'])}", compared_value=f"기준 재계산 {_won(x['expected'])}" + (f" (DET {x['det']}, FTR/RET {x['ftr']})" if t["method"] == "detailed" else " (간이법 평균 가중치)"), reason="표 기재 가중치가 산정 기준과 다릅니다.", suggestion="복잡도·가중치 재산정", question_to_vendor=f"'{x['name']}'의 가중치 산출 근거를 제시해 주십시오.", document_ids=[d.id]))
        seen: dict[tuple[str, str], int] = {}
        for x in t["rows"]:
            k = (x["type"], normalize_name(x["name"]))
            if x["name"] and k in seen:
                out.append(_f(category="FP 중복", verdict="check", location=f"{d.file_name} '{t['sheet']}' {x['row']}행", original_value=f"{x['name']} ({x['type']})", compared_value=f"{seen[k]}행과 같은 이름·유형", reason="같은 단위프로세스가 두 번 계산되었을 가능성(사용자·관리자 화면 중복 포함).", suggestion="처리·데이터 차이가 없으면 1건으로 통합", question_to_vendor=f"{seen[k]}행과 {x['row']}행은 어떤 처리 차이가 있습니까?", document_ids=[d.id]))
            else:
                seen[k] = x["row"]
        if t["total_declared"] is not None:
            if t["total_declared"] != decl_sum:
                out.append(_f(category="FP 합계", verdict="diff", location=f"{d.file_name} {t['total_loc']}", original_value=f"기재 합계 {_won(t['total_declared'])}", compared_value=f"행 합계 {_won(decl_sum)}", reason="FP 합계가 행별 점수의 합과 다릅니다.", suggestion="합계 수식 확인·수정", question_to_vendor="FP 합계 산출 방식을 확인해 주십시오.", document_ids=[d.id]))
            else:
                out.append(_f(category="FP 합계", verdict="match", location=f"{d.file_name} {t['total_loc']}", original_value=_won(t["total_declared"]), compared_value=_won(decl_sum), reason="기재 합계와 행 합계가 같습니다.", document_ids=[d.id]))
        if not wrong:
            out.append(_f(category="FP 가중치", verdict="match", location=d.file_name, reason=f"{len(t['rows'])}행 가중치가 기준({'정통법 복잡도표' if t['method'] == 'detailed' else '간이법'})과 일치(재계산 합 {_won(exp_sum)}).", document_ids=[d.id]))

    # C. 금액 재계산
    cost_docs = by_kind.get("vendor_cost", [])
    if not cost_docs:
        out.append(_f(category="금액 산식", verdict="insufficient", reason="금액 산출서가 등록되지 않아 산식·VAT·합계를 검산하지 못했습니다(오류 아님)."))
    for d in cost_docs:
        vals = parse_cost_values(files.get(d.id), lines[d.id], d.format in ("xlsx", "xlsm"))
        loc = lambda k: f"{d.file_name} {vals[k][1]}"  # noqa: E731
        if not vals:
            out.append(_f(category="금액 산식", verdict="insufficient", location=d.file_name, reason="금액 항목(개발원가·이윤·부가세 등)을 찾지 못했습니다.", document_ids=[d.id]))
            continue
        std_price = Decimal(std["unit_price"]["value"])
        if "unit" in vals:
            u = vals["unit"][0]
            out.append(_f(category="단가", verdict="match" if u == std_price else "check", location=loc("unit"), original_value=f"{_won(u)}원", compared_value=f"기준 {std['id']} {_won(std_price)}원", reason="FP당 단가가 기준과 같습니다." if u == std_price else "FP당 단가가 등록된 기준과 다릅니다(적용 연도·판본 차이 가능).", question_to_vendor="" if u == std_price else "적용한 FP당 단가의 근거(가이드 판본·연도)를 제시해 주십시오.", document_ids=[d.id]))
        if "fp" in vals and fp_total_from_table is not None and vals["fp"][0] != fp_total_from_table:
            out.append(_f(category="문서 간 불일치", verdict="diff", location=loc("fp"), original_value=f"금액 산출서 FP {_won(vals['fp'][0])}", compared_value=f"FP 산정표 합계 {_won(fp_total_from_table)}", reason="금액 산출서의 총 FP와 FP 산정표 합계가 다릅니다.", question_to_vendor="두 문서의 FP 차이 사유를 설명해 주십시오.", document_ids=[x.id for x in cost_docs + fp_docs]))
        factors = [vals[k][0] for k in ("size", "integration", "performance", "compatibility", "security") if k in vals]
        if all(k in vals for k in ("fp", "unit", "dev")) and len(factors) == 5:
            calc = vals["fp"][0] * vals["unit"][0]
            for f in factors:
                calc *= f
            dev = vals["dev"][0]
            ok = abs(calc - dev) <= Decimal(10)
            out.append(_f(category="금액 산식", verdict="match" if ok else "diff", location=loc("dev"), original_value=f"개발원가 {_won(dev)}원", compared_value=f"FP×단가×보정계수 = {_won(calc.quantize(Decimal('0.01')))}원", reason="개발원가가 산식과 일치합니다." if ok else "개발원가가 FP×단가×보정계수 산식과 다릅니다.", question_to_vendor="" if ok else "개발원가 산출 과정을 제시해 주십시오.", document_ids=[d.id]))
        elif "dev" in vals:
            out.append(_f(category="금액 산식", verdict="insufficient", location=loc("dev"), reason="FP·단가·보정계수 5종을 모두 찾지 못해 개발원가 산식을 재계산하지 못했습니다.", document_ids=[d.id]))
        if "dev" in vals and "profit" in vals and vals["dev"][0]:
            rate = vals["profit"][0] / vals["dev"][0]
            out.append(_f(category="이윤", verdict="check", location=loc("profit"), original_value=f"이윤 {_won(vals['profit'][0])}원", compared_value=f"개발원가 대비 {(rate * 100):.2f}%", reason="이윤율 기준 원문을 확인하지 못해 적정성은 판정하지 않습니다(FP_RULES 미확인 항목).", question_to_vendor="적용 이윤율과 근거를 제시해 주십시오.", document_ids=[d.id]))
        if all(k in vals for k in ("dev", "profit", "supply")):
            s = vals["dev"][0] + vals["profit"][0] + (vals["direct"][0] if "direct" in vals else 0)
            ok = abs(s - vals["supply"][0]) <= Decimal(1)
            out.append(_f(category="합계", verdict="match" if ok else "diff", location=loc("supply"), original_value=f"기재 {_won(vals['supply'][0])}원", compared_value=f"개발원가+이윤+직접경비 = {_won(s)}원", reason="부가세 별도 금액 합계가 일치합니다." if ok else "부가세 별도 금액이 구성요소 합과 다릅니다.", question_to_vendor="" if ok else "합계 차이 사유를 설명해 주십시오.", document_ids=[d.id]))
        if "supply" in vals and "vat" in vals:
            exp = vals["supply"][0] * Decimal("0.1")
            ok = abs(exp - vals["vat"][0]) <= Decimal(10)
            out.append(_f(category="부가세", verdict="match" if ok else "diff", location=loc("vat"), original_value=f"부가세 {_won(vals['vat'][0])}원", compared_value=f"공급가액×10% = {_won(exp)}원", reason="부가세가 10%와 일치(끝수 차이 10원 이내)." if ok else "부가세가 공급가액의 10%와 다릅니다.", question_to_vendor="" if ok else "부가세 산출 근거를 확인해 주십시오.", document_ids=[d.id]))
        if all(k in vals for k in ("supply", "vat", "total")):
            s = vals["supply"][0] + vals["vat"][0]
            ok = abs(s - vals["total"][0]) <= Decimal(1)
            out.append(_f(category="합계", verdict="match" if ok else "diff", location=loc("total"), original_value=f"총액 {_won(vals['total'][0])}원", compared_value=f"공급가액+부가세 = {_won(s)}원", reason="총액이 일치합니다." if ok else "총액이 공급가액+부가세와 다릅니다.", question_to_vendor="" if ok else "총액 차이 사유를 설명해 주십시오.", document_ids=[d.id]))
        for loc_, ln in lines[d.id]:
            if EQUIP.search(ln):
                out.append(_f(category="장비 포함 여부", verdict="check", location=f"{d.file_name} {loc_}", original_value=ln[:120], reason="금액 문서에 장비 관련 표현이 있습니다. SW 개발비에 장비 구매비가 섞였는지 확인이 필요합니다.", question_to_vendor="장비 구매비가 개발비에 포함되었습니까?", document_ids=[d.id]))
                break

    # D. 문장 검토(불명확 범위, 연계, 성능·보안)
    vague_n = iface_n = perf_n = 0
    for d in task_docs + by_kind.get("vendor_design", []):
        for loc, ln in lines[d.id]:
            if vague_n < 15 and VAGUE.search(ln) and len(ln) > 10:
                vague_n += 1
                out.append(_f(category="범위 불명확", verdict="check", location=f"{d.file_name} {loc}", original_value=ln[:200], reason=f"범위가 불명확한 표현('{VAGUE.search(ln).group(0).strip()}')이 있어 과업 범위 분쟁 가능성이 있습니다.", suggestion="대상·수량·조건을 구체적으로 기재", question_to_vendor="해당 문장의 구체적 범위(대상·수량)를 명시해 주십시오.", document_ids=[d.id]))
            if iface_n < 10 and IFACE.search(ln) and len(ln) > 10:
                iface_n += 1
                out.append(_f(category="연계 조건", verdict="check", location=f"{d.file_name} {loc}", original_value=ln[:200], reason="연계 가능 여부·방식은 문서만으로 확정할 수 없습니다(상대 시스템 확인 필요).", suggestion="상대 기관 연계 가능 여부·방식·일정 확인서 확보", question_to_vendor="상대 시스템 연계 방식과 협의 상태를 알려 주십시오.", document_ids=[d.id], doc_check_only=False))
            if perf_n < 10 and PERF_SEC.search(ln):
                perf_n += 1
                out.append(_f(category="성능·보안", verdict="check", location=f"{d.file_name} {loc}", original_value=ln[:200], reason="실제 성능·보안 충족은 문서로 확정할 수 없고 시험·점검으로만 확인할 수 있습니다.", suggestion="검수 단계 성능시험·보안점검 결과서 제출 조건 명시", question_to_vendor="해당 요구의 시험·점검 방법을 제시해 주십시오.", document_ids=[d.id], doc_check_only=False))

    # E. 우리 요구사항 반영 여부
    vendor_text = " ".join(ln for d in task_docs + by_kind.get("vendor_design", []) for _, ln in lines[d.id])
    if p.requirements and vendor_text:
        vt = re.sub(r"\s", "", vendor_text)
        hit = 0
        for r in p.requirements:
            if r.kind != "functional" or r.status == "excluded":
                continue
            key = re.sub(r"\s", "", base_noun(r.title))
            if key and key in vt:
                hit += 1
            else:
                out.append(_f(category="요구사항 누락", verdict="diff", location=f"프로젝트 {r.id}", original_value=r.title, compared_value="업체 과업·설계 문서에서 찾지 못함", reason=f"우리 요구사항 '{r.title}'의 핵심어('{base_noun(r.title)}')가 업체 문서에 없습니다(표현 차이 가능).", suggestion="업체 문서에 반영 위치 확인", question_to_vendor=f"'{r.title}' 요구는 어디에 반영되었습니까?", document_ids=[d.id for d in task_docs + by_kind.get('vendor_design', [])]))
        if hit:
            out.append(_f(category="요구사항 반영", verdict="match", reason=f"우리 기능 요구사항 {hit}건의 핵심어가 업체 문서에 나타납니다(핵심어 기준, 세부 내용은 별도 확인)."))
    elif p.requirements and not vendor_text:
        out.append(_f(category="요구사항 반영", verdict="insufficient", reason="업체 과업·설계 문서 텍스트가 없어 우리 요구사항 반영 여부를 확인하지 못했습니다."))
    return out
