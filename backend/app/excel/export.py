"""프로젝트 → 엑셀(.xlsx) 산출물. 산출내역서처럼 읽고 인쇄하기 쉬운 업무용 서식(2026-10-30 개편, docs/EXCEL_FORMAT.md).

- 모든 사용자 텍스트는 문자열 셀로 저장(수식 실행 방지).
- 금액·FP는 Excel 수식으로 두고, Python 계산값은 '검산' 시트에 함께 적어 수식값과 비교한다. openpyxl은 수식 캐시값을
  쓰지 못하므로 통합문서에 '열 때 전체 재계산'을 설정한다.
- 산식·기준값·끝수 규칙은 계산 엔진(fp/cost.py, fp/engine.py)과 같다. 서식 개편에서 바꾸지 않는다.
- 인쇄: A4, 눈금선 숨김, 가로는 한 페이지 폭에 맞추고 세로는 자연스럽게 나눔, 머리행 반복, 페이지 번호, 실제 데이터까지만 인쇄 영역.
"""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.worksheet import Worksheet

from ..fp import engine
from ..fp.cost import PHASE_KEYS, PHASE_LABELS
from ..fp.estimates import estimate_state
from ..models import EstimateVersion, Project
from ..services import checklist as cl
from ..services.taskdraft import build_sections
from . import labels as L

# ---------------------------------------------------------------- 서식(참고: 산출내역서 FP금액·FP산정 시트의 업무용 표 스타일)
FONT_NAME = "맑은 고딕"
F_BODY = Font(name=FONT_NAME, size=10)
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_TITLE = Font(name=FONT_NAME, size=15, bold=True)
F_SUB = Font(name=FONT_NAME, size=9, color="595959")
F_SEC = Font(name=FONT_NAME, size=11, bold=True)
F_UNIT = Font(name=FONT_NAME, size=8, color="595959")
F_TOTAL = Font(name=FONT_NAME, size=11, bold=True)
F_WARN = Font(name=FONT_NAME, size=10, bold=True, color="C00000")
F_REF = Font(name=FONT_NAME, size=10, bold=True, color="1F4E79")
FILL_HEAD = PatternFill("solid", fgColor="FFF2CC")  # 머리글·구분(연한 노랑)
FILL_LABEL = PatternFill("solid", fgColor="F2F2F2")  # 항목명
FILL_TOTAL = PatternFill("solid", fgColor="D9D9D9")  # 소계·합계·총액
FILL_REF = PatternFill("solid", fgColor="DDEBF7")  # 참고값(목표예산 시나리오)
FILL_WARN = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="808080")
MEDIUM = Side(style="medium", color="404040")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NO_BORDER = Border()
WON = '#,##0'
WON_DEC = '#,##0.00'
FP_FMT = '#,##0.0'
RATIO = '0.0000'
AL = {
    "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
    "top": Alignment(horizontal="left", vertical="top", wrap_text=True),
    "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
    "right": Alignment(horizontal="right", vertical="center"),
}

SHEET_NAMES = [
    "사업개요", "개발비_기준", "목표예산_시나리오", "FP_간이법", "FP_정통법", "참조_보정계수",
    "요구사항", "기능설계", "화면_데이터_연계", "과업내용", "검토의견", "HW_SW추천", "근거_변경이력", "검산",
]

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def put(ws: Worksheet, row: int, col: int, value: Any, *, fmt: str | None = None, bold: bool = False, formula: bool = False,
        al: str | None = None, fill: PatternFill | None = None, font: Font | None = None, border: bool = True):
    """셀 기록. formula=True일 때만 '='로 시작하는 값을 수식으로 둔다. 그 외 문자열은 항상 텍스트.
    정렬 기본값: 숫자·수식은 오른쪽, 문자는 왼쪽(줄바꿈)."""
    c = ws.cell(row=row, column=col)
    if value is None:
        value = ""
    if isinstance(value, Decimal):
        value = float(value)
    if isinstance(value, bool):
        value = "예" if value else "아니오"
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(x) for x in value)
    c.value = value
    if isinstance(value, str) and not (formula and value.startswith("=")):
        c.data_type = "s"  # 수식 실행 방지
    if fmt:
        c.number_format = fmt
    if al is None:
        al = "right" if (isinstance(value, (int, float)) or (formula and fmt)) else "left"
    c.alignment = AL[al]
    c.border = BORDER if border else NO_BORDER
    c.font = font or (F_BOLD if bold else F_BODY)
    if fill:
        c.fill = fill
    return c


def widths(ws: Worksheet, ws_widths: list[float]) -> None:
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def header(ws: Worksheet, row: int, titles: list[str], col_widths: list[float] | None = None, start: int = 1):
    for i, t in enumerate(titles, start=start):
        put(ws, row, i, t, bold=True, al="center", fill=FILL_HEAD)
    ws.row_dimensions[row].height = 30
    if col_widths:
        widths(ws, col_widths)


def title(ws: Worksheet, text: str, sub: str = "", ncols: int = 4, unit: str = "") -> int:
    """제목(1행, 병합)·부제(2행). 반환: 다음 내용 시작 행."""
    ws.sheet_view.showGridLines = False
    end = get_column_letter(max(ncols, 1))
    c = ws.cell(row=1, column=1, value=text)
    c.data_type, c.font, c.alignment = "s", F_TITLE, Alignment(horizontal="left", vertical="center", wrap_text=True)
    if ncols > 1:
        ws.merge_cells(f"A1:{end}1")
    ws.row_dimensions[1].height = 36 if len(text) < 60 else 54
    if sub:
        s = ws.cell(row=2, column=1, value=sub)
        s.data_type, s.font, s.alignment = "s", F_SUB, Alignment(horizontal="left", vertical="center", wrap_text=True)
        if ncols > 1:
            ws.merge_cells(f"A2:{end}2")
        ws.row_dimensions[2].height = 15 * max(1, min(4, len(sub) // 110 + 1))
    if unit:
        u = ws.cell(row=3, column=max(ncols, 1), value=unit)
        u.data_type, u.font, u.alignment = "s", F_UNIT, Alignment(horizontal="right")
    return 4


def section(ws: Worksheet, row: int, text: str, unit: str = "", ncols: int = 0) -> int:
    c = ws.cell(row=row, column=1, value=f"○ {text}")
    c.data_type, c.font = "s", F_SEC
    if unit and ncols:
        u = ws.cell(row=row, column=ncols, value=unit)
        u.data_type, u.font, u.alignment = "s", F_UNIT, Alignment(horizontal="right")
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws: Worksheet, row: int, text: str, ncols: int, font: Font = F_SUB, fill: PatternFill | None = None) -> int:
    c = ws.cell(row=row, column=1, value=text)
    c.data_type, c.font, c.alignment = "s", font, Alignment(horizontal="left", vertical="center", wrap_text=True)
    if fill:
        c.fill = fill
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = 15 * max(1, min(6, len(text) // max(20, int(sum((ws.column_dimensions[get_column_letter(i)].width or 9) for i in range(1, ncols + 1)) * 0.9)) + 1))
    return row + 1


def page_setup(ws: Worksheet, landscape: bool = True, ncols: int | None = None, title_rows: str | None = None):
    """A4. 가로 폭만 한 페이지에 맞추고(세로는 자연스럽게 나눔) 페이지 번호·머리행 반복·실제 데이터 인쇄 영역."""
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.45
    ws.page_margins.top, ws.page_margins.bottom = 0.6, 0.6
    ws.page_margins.header, ws.page_margins.footer = 0.25, 0.3
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = "&A"
    ws.oddFooter.right.size = 8
    if title_rows:
        ws.print_title_rows = title_rows
    last_col = get_column_letter(ncols or ws.max_column)
    ws.print_area = f"A1:{last_col}{ws.max_row}"


def finish_table(ws: Worksheet, header_row: int, last_row: int, ncols: int, landscape: bool = True, freeze_col: int | None = None, header_top: int | None = None):
    ws.freeze_panes = f"{get_column_letter(freeze_col or (2 if ncols > 3 else 1))}{header_row + 1}"
    if last_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last_row}"
    page_setup(ws, landscape, ncols, f"{header_top or header_row}:{header_row}")


def not_run(ws: Worksheet, msg: str):
    widths(ws, [100])
    r = title(ws, ws.title.replace("_", " "), "미실행 — 이 시트의 내용은 아직 만들어지지 않았습니다('문제 없음'이 아님).", 1)
    put(ws, r, 1, f"미실행: {msg}", font=F_WARN, fill=FILL_WARN)
    page_setup(ws, False, 1)


def _mode_text(p: Project) -> str:
    origins: dict[str, int] = {}
    for coll in ("requirements", "functions", "fp_items", "screens", "data_entities", "interfaces"):
        for it in getattr(p, coll):
            origins[it.origin] = origins.get(it.origin, 0) + 1
    return ", ".join(f"{L.ORIGIN.get(k, k)} {v}건" for k, v in sorted(origins.items())) or "항목 없음"


def pick_estimate(p: Project, estimate_id: str | None) -> EstimateVersion | None:
    if estimate_id:
        return next((e for e in p.estimates if e.id == estimate_id), None)
    return p.estimates[-1] if p.estimates else None


def build_workbook(p: Project, std: dict[str, Any], changes: list[dict[str, Any]], mode_label: str, estimate_id: str | None = None) -> tuple[Workbook, dict[str, Any]]:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    sheets = {name: wb.create_sheet(name) for name in SHEET_NAMES}
    fp_sum = engine.summary(p, std)
    est = pick_estimate(p, estimate_id)
    checks: list[dict[str, Any]] = []  # 검산 시트용: 수식 셀과 Python 계산값

    _overview(sheets["사업개요"], p, fp_sum, est, mode_label)
    if est:
        cost_cells = _cost(sheets["개발비_기준"], p, est, checks)
        _scenarios(sheets["목표예산_시나리오"], est, cost_cells, checks)
    else:
        not_run(sheets["개발비_기준"], "기준 산정(견적 버전)을 아직 만들지 않았습니다. 'FP·개발비' 단계에서 기준 산정을 실행하세요(설정 저장만으로는 견적이 만들어지지 않음).")
        not_run(sheets["목표예산_시나리오"], "기준 산정이 없어 목표예산 시나리오를 만들 수 없습니다.")
    _fp_simple(sheets["FP_간이법"], p, std, fp_sum, checks)
    _fp_detailed(sheets["FP_정통법"], p, std, fp_sum, checks)
    _factor_ref(sheets["참조_보정계수"], std)
    _requirements(sheets["요구사항"], p, fp_sum)
    _functions(sheets["기능설계"], p)
    _design_lists(sheets["화면_데이터_연계"], p)
    _task(sheets["과업내용"], p)
    _reviews(sheets["검토의견"], p)
    _tech(sheets["HW_SW추천"], p)
    _trace(sheets["근거_변경이력"], p, changes, fp_sum)
    _checks(sheets["검산"], checks)
    meta = {"estimate_id": est.id if est else None, "checks": checks, "fp_summary": {"simple": fp_sum["simple"]["total_fp"], "detailed": fp_sum["detailed"]["total_fp"]}}
    return wb, meta


def to_bytes(wb: Workbook) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------- 사업개요(첫 화면 요약)
def _overview(ws, p: Project, fp_sum, est, mode_label: str):
    ov = p.overview
    widths(ws, [30, 74])
    r = title(ws, p.name, f"정보화 설계 도우미 산출물 · 생성 모드: {mode_label} · 발주자 검토용 초안", 2)
    r = section(ws, r, "산정 요약")
    m = est.result.get("money", {}) if est else {}
    basis = (est.result.get("basis") or {}) if est else {}
    state = estimate_state(p, est) if est else ""
    if est and m.get("complete"):
        money_txt = f"{m['total']:,}원 (부가세 포함) · 공급가액 {m['supply']:,}원"
    elif est:
        money_txt = "미산정(금액 없음) — " + ("집계된 FP 0건" if basis.get("basis") == "empty" else "견적 미완료")
    else:
        money_txt = "기준 산정 미실행"
    scen = est.scenarios[-1] if est and est.scenarios else None
    summary = [
        ("산정 상태", f"간이법 {_basis_text(fp_sum, 'simple')} · 정통법 {_basis_text(fp_sum, 'detailed')}"),
        ("산정 기준일(견적 버전)", f"{est.created_at[:10]} · {est.id} {est.label}".strip() + (" · 현재 데이터와 일치" if state == "current" else " · 갱신 필요(이후 변경 있음)") if est else "기준 산정 없음"),
        ("적용 방법·기준", f"{L.METHOD[est.method]} · {est.standard_snapshot.get('id', '')}" if est else f"{L.METHOD[p.cost_settings.method]}(설정) · {p.cost_settings.standard_id}"),
        ("기준 FP", f"{m['total_fp']} FP ({basis.get('label', '이전 형식 견적')})" if est and m.get("total_fp") is not None else f"현재 집계 {fp_sum['simple']['total_fp']} FP(간이법) — 견적 아님"),
        ("개발비(기준 산정)", money_txt),
        ("목표예산", (f"목표 {ov.target_amount:,}원 ({'부가세 포함' if ov.target_includes_vat else '부가세 별도'})" if ov.budget_mode == "target" and ov.target_amount else "지정 안 함")
         + (f" · 참고 시나리오 {scen.id}: {scen.result['total']:,}원(조정률 {scen.result.get('adjustment_ratio_display', '')}, 공식 산정 아님)" if scen else "")),
        ("확인 목록", ", ".join(f"{k} {v}건" for k, v in cl.summarize(cl.checklist(p, fp_sum)).items() if v) or "남은 항목 없음"),
    ]
    for k, v in summary:
        put(ws, r, 1, k, bold=True, fill=FILL_LABEL)
        hl = k in ("개발비(기준 산정)", "기준 FP")
        put(ws, r, 2, v, font=F_TOTAL if hl else None, fill=FILL_TOTAL if hl and est and m.get("complete") else None)
        r += 1
    if basis.get("basis") in ("partial", "estimated"):
        r = note(ws, r, f"※ 산정 상태 '{basis['label']}': " + "; ".join(basis.get("reasons", [])) + " — FP 전문가 검토 전의 확정 견적이 아닙니다.", 2, F_WARN)
    r += 1
    r = section(ws, r, "사업 정보")
    phases = "·".join(PHASE_LABELS[k] for k in PHASE_KEYS if getattr(ov.phases, k)) or "없음"
    rows = [
        ("프로젝트 ID", p.id),
        ("사업 유형", L.BUSINESS_TYPE[ov.business_type]),
        ("시스템 형태", ", ".join(L.SYSTEM_KINDS[k] for k in ov.system_kinds) or "미정"),
        ("시작 방식", L.START_MODE[ov.start_mode]),
        ("처음 구상(원문)", ov.idea_text or "미입력"),
        ("목적", ov.purpose or "미입력"),
        ("추진 배경", ov.background or "미입력"),
        ("이용자", ov.users or "미입력"),
        ("구축 방식", ov.build_method or "미입력"),
        ("사업 기간(개월)", ov.period_months if ov.period_months else "미정"),
        ("금액 조건", L.BUDGET_MODE[ov.budget_mode]),
        ("포함 단계", phases),
        ("기관 환경", ov.org_environment or "미입력"),
        ("FP 측정 대상(애플리케이션 경계)", ov.boundary_name or "미지정(경계 확인 필요)"),
        ("장비 구매비", "개발비 합계에서 제외(HW·SW 추천은 설계 보조 정보)"),
        ("항목 출처 구분", _mode_text(p)),
        ("집계 FP(간이법/정통법)", f"{fp_sum['simple']['total_fp']} / {fp_sum['detailed']['total_fp']} (확정 {fp_sum['counts']['confirmed_counted']}건 + 업무 확인 {fp_sum['counts']['scope_counted']}건)"),
        ("업무 확인 전 FP 후보 수", fp_sum["counts"]["awaiting_scope"]),
        ("미분류 FP 항목 수", fp_sum["counts"]["unclassified"]),
        ("미답변 질문 수", sum(1 for q in p.questions if q.status == "open")),
        ("재검토 필요 항목 수", sum(1 for c in ("requirements", "functions", "fp_items") for it in getattr(p, c) if it.needs_review)),
        ("데이터 버전(최종 수정)", p.updated_at),
    ]
    for k, v in rows:
        put(ws, r, 1, k, bold=True, fill=FILL_LABEL)
        put(ws, r, 2, v, al="left")
        r += 1
    r += 1
    r = section(ws, r, "견적 버전 목록")
    for i, t in enumerate(["견적ID · 이름 · 생성일", "방법 · FP · 총액 · 상태"], 1):
        put(ws, r, i, t, bold=True, al="center", fill=FILL_HEAD)
    r += 1
    if not p.estimates:
        put(ws, r, 1, "-")
        put(ws, r, 2, "기준 산정 없음")
        r += 1
    for e in reversed(p.estimates):
        em = e.result.get("money", {})
        eb = (e.result.get("basis") or {}).get("label", "")
        put(ws, r, 1, f"{e.id} {e.label} · {e.created_at[:16].replace('T', ' ')}" + (" (이 파일의 기준)" if est and e.id == est.id else ""), bold=bool(est and e.id == est.id))
        put(ws, r, 2, f"{L.METHOD[e.method]} · {em.get('total_fp', '-')} FP · " + (f"{em['total']:,}원" if em.get("complete") else "금액 없음") + f" · {eb} · {'현재 데이터와 일치' if estimate_state(p, e) == 'current' else '갱신 필요'}" + (f" · 시나리오 {len(e.scenarios)}건" if e.scenarios else ""))
        r += 1
    r += 1
    r = section(ws, r, "가정 사항")
    for i, t in enumerate(["ID", "가정 내용"], 1):
        put(ws, r, i, t, bold=True, al="center", fill=FILL_HEAD)
    r += 1
    if not p.assumptions:
        put(ws, r, 1, "-")
        put(ws, r, 2, "기록된 가정 없음")
    for a in p.assumptions:
        put(ws, r, 1, a.id, al="center")
        put(ws, r, 2, f"{a.text} (이유: {a.reason})" if a.reason else a.text)
        r += 1
    page_setup(ws, False, 2)


def _basis_text(fp_sum, method: str) -> str:
    b = fp_sum["basis"][method]
    total = fp_sum[method]["total_fp"]
    return f"{b['label']}" + (f" {total} FP" if b["basis"] != "empty" else "(집계 대상 0건)")


def _status_text(it, row) -> str:
    if it.status == "confirmed":
        return "확정(기술 검토)"
    if row.get("count_basis") == "scope":
        return "업무 확인(추정)"
    return L.STATUS[it.status] + (" · 업무 확인" if it.scope_confirmed else "")


# ---------------------------------------------------------------- 개발비(산출내역서 'FP금액' 구성)
def _round_formula(expr: str, unit: int, mode: str) -> str:
    if mode == "down":
        return f"ROUNDDOWN(({expr})/{unit},0)*{unit}"
    return f"ROUND(({expr})/{unit},0)*{unit}"


COST_COLS = 10  # A~J


def _cost(ws, p: Project, est: EstimateVersion, checks) -> dict[str, str]:
    m = est.result["money"]
    std = est.standard_snapshot
    state = estimate_state(p, est)
    basis = est.result.get("basis") or {}
    widths(ws, [12, 11, 8, 11, 8, 10, 8, 9, 14, 14])
    sub = (f"기준: {std.get('id')}({std.get('title', '')}) · 방법: {L.METHOD[est.method]} · 견적 {est.id} '{est.label}' {est.created_at[:16].replace('T', ' ')} · "
           + ("현재 데이터와 일치" if state == "current" else "※ 이후 FP·설정이 바뀜: 갱신 필요"))
    r = title(ws, f"'{p.name}' 기능점수 기반 소프트웨어 개발비 산정(기준 산정)", sub, COST_COLS)
    cells: dict[str, str] = {}
    if not m.get("complete"):
        msg = "미산정(금액 없음)" if basis.get("basis") == "empty" or Decimal(str(m.get("total_fp", "0"))) == 0 else "미완료 견적(금액 없음)"
        r = note(ws, r, f"{msg} — " + " ".join(est.warnings), COST_COLS, F_WARN, FILL_WARN)
        page_setup(ws, False, COST_COLS)
        return cells
    if basis.get("basis") in ("partial", "estimated"):
        r = note(ws, r, f"산정 상태: {basis.get('label')} — " + "; ".join(basis.get("reasons", [])) + " — FP 전문가 검토 전의 확정 견적이 아닙니다.", COST_COLS, F_WARN, FILL_WARN)
    unit, mode = m["rounding"]["unit"], m["rounding"]["mode"]
    by_type = (est.result.get("fp") or {}).get("by_type", {})
    types = ("ILF", "EIF", "EI", "EO", "EQ")
    # 구획 행 위치를 먼저 정한다(개발원가 표가 아래 구획 셀을 참조)
    fx = m["factors"]
    direct_items = m["direct_expenses"] or [{"name": "-", "basis": "직접경비 없음(장비 구매비 제외)", "amount": 0}]
    r_cost_h = r + 1  # 개발원가 산정 머리글 2행
    r_cost_v = r_cost_h + 2
    r_sum = r_cost_v + 2  # 금액 요약 시작
    sum_rows = ["dev", "profit", "direct", "supply", "vat", "total"]
    r_direct = r_sum + len(sum_rows) + 2  # 직접경비 구획 제목
    r_direct_h = r_direct + 1
    r_direct_total = r_direct_h + 1 + len(direct_items)
    r_factor = r_direct_total + 2
    r_factor_h = r_factor + 1
    r_size = r_factor_h + 1
    r_fac0 = r_size + 1
    r_phase = r_fac0 + len(fx)
    r_fp = r_phase + 2
    r_fp_h = r_fp + 1

    # ○ 개발원가 산정
    section(ws, r, "개발원가 산정", "(단위: 원)", COST_COLS)
    heads = [("총기능점수", 1, 1), ("기능점수당 단가", 2, 2), ("보정계수", 3, 8), ("개발원가(끝수 처리 전)", 9, 9), ("개발원가", 10, 10)]
    for text, c1, c2 in heads:
        put(ws, r_cost_h, c1, text, bold=True, al="center", fill=FILL_HEAD)
        for cc in range(c1 + 1, c2 + 1):
            put(ws, r_cost_h, cc, "", fill=FILL_HEAD)
        if c2 > c1:
            ws.merge_cells(start_row=r_cost_h, start_column=c1, end_row=r_cost_h, end_column=c2)
        else:
            put(ws, r_cost_h + 1, c1, "", fill=FILL_HEAD)
            ws.merge_cells(start_row=r_cost_h, start_column=c1, end_row=r_cost_h + 1, end_column=c1)
    for i, t in enumerate(["규모", "연계복잡성", "성능", "운영환경", "보안성", "단계비율"], 3):
        put(ws, r_cost_h + 1, i, t, bold=True, al="center", fill=FILL_HEAD)
    ws.row_dimensions[r_cost_h].height = 20
    ws.row_dimensions[r_cost_h + 1].height = 20
    fp_cell = f"A{r_cost_v}"
    put(ws, r_cost_v, 1, f"=SUM(B{r_fp_h + 2}:F{r_fp_h + 2})", fmt=FP_FMT, formula=True, al="center")
    put(ws, r_cost_v, 2, m["unit_price"], fmt=WON, al="center")
    put(ws, r_cost_v, 3, f"=J{r_size}", fmt=RATIO, formula=True, al="center")
    for i in range(len(fx)):
        put(ws, r_cost_v, 4 + i, f"=J{r_fac0 + i}", fmt="0.00", formula=True, al="center")
    put(ws, r_cost_v, 8, f"=J{r_phase}", fmt="0.00", formula=True, al="center")
    factor_refs = "*".join(f"{get_column_letter(c)}{r_cost_v}" for c in range(3, 9))
    raw = f"I{r_cost_v}"
    put(ws, r_cost_v, 9, f"=A{r_cost_v}*B{r_cost_v}*{factor_refs}", fmt=WON_DEC, formula=True)
    dev = f"J{r_cost_v}"
    put(ws, r_cost_v, 10, "=" + _round_formula(raw, unit, mode), fmt=WON, formula=True)
    ws.row_dimensions[r_cost_v].height = 30
    cells |= {"fp": fp_cell, "price": f"B{r_cost_v}", "size": f"C{r_cost_v}", "dev_raw": raw, "dev": dev}
    note(ws, r_cost_v + 1, f"개발원가 = 총기능점수 × 기능점수당 단가 × 보정계수(규모·연계·성능·운영환경·보안) × 단계비율 · 끝수 처리: {m['rounding']['label']}(시스템 설정, 공식 규정 미확인) · 단가 근거: {std['unit_price']['basis']}"
         + (" / 사용자 입력 단가" if m.get("unit_price_source") == "user" else ""), COST_COLS)

    # 금액 요약(참고 파일: 개발원가 → 이윤 → 직접경비 → 개발비(부가세 별도) → 부가세 → 총액)
    pr_cell = f"I{r_sum + 1}"
    vr_cell = f"I{r_sum + 4}"
    lines = {
        "dev": ("합계(보정 후 개발원가)", None, None, f"={dev}"),
        "profit": ("이윤", Decimal(m["profit_rate"]), "0%", "=" + _round_formula(f"J{r_sum}*{pr_cell}", unit, mode)),
        "direct": ("직접경비", None, None, f"=J{r_direct_total}"),
        "supply": ("소프트웨어 개발비(부가세 별도)", None, None, f"=SUM(J{r_sum}:J{r_sum + 2})"),
        "vat": ("부가가치세", Decimal(m["vat_rate"]), "0%", "=" + _round_formula(f"J{r_sum + 3}*{vr_cell}", unit, mode)),
        "total": ("총액(부가세 포함)", None, None, f"=J{r_sum + 3}+J{r_sum + 4}"),
    }
    for i, key in enumerate(sum_rows):
        rr = r_sum + i
        label, rate, rfmt, formula = lines[key]
        strong = key in ("supply", "total")
        put(ws, rr, 1, label, bold=True, al="center", fill=FILL_TOTAL if strong else FILL_LABEL, font=F_TOTAL if key == "total" else None)
        for cc in range(2, 9):
            put(ws, rr, cc, "", fill=FILL_TOTAL if strong else FILL_LABEL)
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=8)
        put(ws, rr, 9, rate if rate is not None else "", fmt=rfmt, al="center", fill=FILL_HEAD if rate is not None else (FILL_TOTAL if strong else None))
        put(ws, rr, 10, formula, fmt=WON, formula=True, bold=strong, fill=FILL_TOTAL if strong else None, font=F_TOTAL if key == "total" else None)
        ws.row_dimensions[rr].height = 22
        cells[key] = f"J{rr}"
    cells["profit_rate"] = pr_cell
    cells["vat_rate"] = vr_cell
    note(ws, r_sum + len(sum_rows), f"이윤율 근거: {std['profit_rate']['basis']} · 부가세율 근거: {std['vat_rate']['basis']} · 장비 구매비·별도 유료 SW는 포함하지 않음", COST_COLS)

    # ○ 직접경비
    section(ws, r_direct, "직접경비", "(단위: 원)", COST_COLS)
    for c1, c2, t in ((1, 2, "구분"), (3, 9, "산출내역"), (10, 10, "금액")):
        put(ws, r_direct_h, c1, t, bold=True, al="center", fill=FILL_HEAD)
        for cc in range(c1 + 1, c2 + 1):
            put(ws, r_direct_h, cc, "", fill=FILL_HEAD)
        if c2 > c1:
            ws.merge_cells(start_row=r_direct_h, start_column=c1, end_row=r_direct_h, end_column=c2)
    for i, d in enumerate(direct_items):
        rr = r_direct_h + 1 + i
        put(ws, rr, 1, d["name"] or "-", al="center")
        put(ws, rr, 2, "")
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
        put(ws, rr, 3, d["basis"] or "-")
        for cc in range(4, 10):
            put(ws, rr, cc, "")
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=9)
        put(ws, rr, 10, d["amount"], fmt=WON)
    put(ws, r_direct_total, 1, "합 계", bold=True, al="center", fill=FILL_TOTAL)
    for cc in range(2, 10):
        put(ws, r_direct_total, cc, "", fill=FILL_TOTAL)
    ws.merge_cells(start_row=r_direct_total, start_column=1, end_row=r_direct_total, end_column=9)
    put(ws, r_direct_total, 10, f"=SUM(J{r_direct_h + 1}:J{r_direct_total - 1})", fmt=WON, formula=True, bold=True, fill=FILL_TOTAL)

    # ○ 보정계수 산정
    section(ws, r_factor, "보정계수 산정")
    for c1, c2, t in ((1, 2, "보정계수 구분"), (3, 9, "복잡도 및 난이도 수준"), (10, 10, "보정계수")):
        put(ws, r_factor_h, c1, t, bold=True, al="center", fill=FILL_HEAD)
        for cc in range(c1 + 1, c2 + 1):
            put(ws, r_factor_h, cc, "", fill=FILL_HEAD)
        if c2 > c1:
            ws.merge_cells(start_row=r_factor_h, start_column=c1, end_row=r_factor_h, end_column=c2)
    frows = [("SW 규모", "= 0.4057 × (ln(총기능점수) − 7.1978)² + 0.8878 (500FP 미만 1.28, 3,000FP 초과 1.153)",
              f"=IF({fp_cell}<500,1.28,IF({fp_cell}>3000,1.153,0.4057*(LN({fp_cell})-7.1978)^2+0.8878))", RATIO)]
    frows += [(f["label"], f["level_text"] if str(f["level_text"]).startswith(f"{f['level']}.") else f"{f['level']}. {f['level_text']}", Decimal(f["value"]), "0.00") for f in fx]
    frows += [("단계 비율", m["phase_text"], Decimal(m["phase_multiplier"]), "0.00")]
    for i, (lab, txt, val, vf) in enumerate(frows):
        rr = r_size + i
        put(ws, rr, 1, lab, bold=True, al="center", fill=FILL_LABEL)
        put(ws, rr, 2, "", fill=FILL_LABEL)
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
        put(ws, rr, 3, txt)
        for cc in range(4, 10):
            put(ws, rr, cc, "")
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=9)
        put(ws, rr, 10, val, fmt=vf, formula=isinstance(val, str), al="center")
        ws.row_dimensions[rr].height = 30
    for i, f in enumerate(fx):
        cells[f["key"]] = f"J{r_fac0 + i}"
    cells["phase"] = f"J{r_phase}"

    # ○ FP 집계(견적 시점 스냅샷 기준)
    section(ws, r_fp, f"FP 집계({L.METHOD[est.method]}, 견적 시점 집계 항목 {len(est.fp_snapshot)}건)")
    for i, t in enumerate(["구분", *types, "계"], 1):
        put(ws, r_fp_h, i, t, bold=True, al="center", fill=FILL_HEAD)
    put(ws, r_fp_h + 1, 1, "기능수", bold=True, al="center", fill=FILL_LABEL)
    put(ws, r_fp_h + 2, 1, "기능점수", bold=True, al="center", fill=FILL_LABEL)
    put(ws, r_fp_h + 3, 1, "비중", bold=True, al="center", fill=FILL_LABEL)
    for i, t in enumerate(types, 2):
        bt = by_type.get(t, {"count": 0, "fp": 0})
        col = get_column_letter(i)
        put(ws, r_fp_h + 1, i, bt["count"], fmt="0", al="center")
        put(ws, r_fp_h + 2, i, Decimal(str(bt["fp"])), fmt=FP_FMT, al="center")
        put(ws, r_fp_h + 3, i, f"=IF($G${r_fp_h + 2}=0,0,{col}{r_fp_h + 2}/$G${r_fp_h + 2})", fmt="0.0%", formula=True, al="center")
    put(ws, r_fp_h + 1, 7, f"=SUM(B{r_fp_h + 1}:F{r_fp_h + 1})", fmt="0", formula=True, al="center", bold=True, fill=FILL_TOTAL)
    put(ws, r_fp_h + 2, 7, f"=SUM(B{r_fp_h + 2}:F{r_fp_h + 2})", fmt=FP_FMT, formula=True, al="center", bold=True, fill=FILL_TOTAL)
    put(ws, r_fp_h + 3, 7, f"=SUM(B{r_fp_h + 3}:F{r_fp_h + 3})", fmt="0.0%", formula=True, al="center", bold=True, fill=FILL_TOTAL)
    r = r_fp_h + 5

    # 검산(수식 셀 ↔ Python 계산값)
    py = {"fp": Decimal(m["total_fp"]), "size": Decimal(m["size_factor"]), "dev_raw": Decimal(m["dev_cost_raw"]), "dev": m["dev_cost"], "profit": m["profit"],
          "direct": m["direct_total"], "supply": m["supply"], "vat": m["vat"], "total": m["total"]}
    for key, label in (("fp", "총 기능점수"), ("size", "규모 보정계수"), ("dev_raw", "개발원가(끝수 전)"), ("dev", "개발원가"), ("profit", "이윤"), ("direct", "직접경비"),
                       ("supply", "SW개발비(VAT 별도)"), ("vat", "부가세"), ("total", "총액")):
        checks.append({"label": f"개발비 {label}", "sheet": ws.title, "formula_cell": cells[key], "python": float(py[key]), "fmt": WON_DEC if key in ("dev_raw", "size") else (FP_FMT if key == "fp" else WON)})

    r = section(ws, r, "계산 근거(Python 계산 과정)")
    for line in m["basis"]:
        r = note(ws, r, line, COST_COLS, F_BODY)
    if est.warnings:
        r += 1
        r = section(ws, r, "주의 사항")
        for w in est.warnings:
            r = note(ws, r, "※ " + w, COST_COLS, F_WARN)
    page_setup(ws, False, COST_COLS)
    return cells


def _scenarios(ws, est: EstimateVersion, cost_cells: dict[str, str], checks):
    if not est.scenarios:
        not_run(ws, f"견적 {est.id}에 목표예산 시나리오가 없습니다(예산 미지정 또는 미실행).")
        return
    widths(ws, [34, 18, 70])
    r = title(ws, "목표예산 시나리오(참고 견적 — 기준 산정 아님)", "공식 측정 FP가 아닌 예산 대응 참고값 · 기준 FP·분류·가중치·요구 기능은 변경하지 않음 · 정식 금액은 '개발비_기준' 시트", 3)
    r = note(ws, r, "※ 이 시트의 조정률·환산 FP·환산 단가는 목표 금액에 맞춘 참고 계산입니다. 공식 산정값으로 사용하지 마세요.", 3, F_REF, FILL_REF)
    base = "'개발비_기준'!"
    for sc in est.scenarios:
        res = sc.result
        r += 1
        r = section(ws, r, f"{sc.id} ({sc.created_at[:16].replace('T', ' ')}) — {res['method_text']}")
        header(ws, r, ["구분", "값", "설명"])
        r += 1

        def row(label, formula, py, fmt, note_txt="", strong=False):
            nonlocal r
            put(ws, r, 1, label, bold=True, fill=FILL_TOTAL if strong else FILL_LABEL)
            put(ws, r, 2, formula if formula else py, fmt=fmt, formula=bool(formula), fill=FILL_REF if strong else None, bold=strong)
            put(ws, r, 3, note_txt)
            r += 1
            return f"B{r - 1}"

        row("목표 금액", None, res["target_amount"], WON, "부가세 포함" if res["target_includes_vat"] else "부가세 별도")
        row("기준 산정 총액", f"={base}{cost_cells['total']}" if "total" in cost_cells else None, res["base_total"], WON, f"견적 {est.id}(정식 기준 산정)")
        row("차이(시나리오 총액 − 기준 총액)", None, res["difference_total"], WON)
        sup = row("목표 공급가액(부가세 별도)", None, res["supply"], WON, "부가세 포함 목표액에서 역산" if res["target_includes_vat"] else "")
        dirc = row("직접경비", None, res["direct_total"], WON, "기준과 동일")
        ratio = row("조정률 r", f"=(({sup}-{dirc})/(1+{base}{cost_cells['profit_rate']}))/{base}{cost_cells['dev_raw']}", Decimal(res["adjustment_ratio"]), "0.000000", "(목표 공급가액 − 직접경비) ÷ (1+이윤율) ÷ 기준 개발원가(끝수 전)")
        row("기준 FP(공식 측정, 변경 없음)", f"={base}{cost_cells['fp']}", Decimal(res["base_total_fp"]), FP_FMT, "기준 FP 보존")
        row("예산 대응 환산 FP(참고값)", f"={base}{cost_cells['fp']}*{ratio}", Decimal(res["equivalent_fp"]), "0.00", "공식 측정 FP 아님")
        row("예산 대응 환산 FP당 단가(참고값)", f"={base}{cost_cells['price']}*{ratio}", Decimal(res["equivalent_unit_price"]), WON_DEC, "공식 단가 아님")
        row("조정 개발원가", None, res["dev_cost"], WON)
        row("조정 이윤", None, res["profit"], WON)
        row("끝수조정(원 단위 일치용)", None, res["rounding_adjustment"], WON, "가상의 비용 항목이 아니라 원 단위 합계를 맞추기 위한 표시")
        row("공급가액 합계", None, res["supply"], WON)
        vat = row("부가가치세", None, res["vat"], WON)
        tot = row("최종 참고 금액(부가세 포함)", f"={sup}+{vat}", res["total"], WON, ("목표액과 일치" if res["matches_target"] else f"목표액과 {res['vat_residual']:+,}원 차이(부가세 끝수)"), strong=True)
        checks.append({"label": f"{sc.id} 조정률", "sheet": ws.title, "formula_cell": ratio, "python": float(Decimal(res["adjustment_ratio"])), "fmt": "0.000000"})
        checks.append({"label": f"{sc.id} 최종 참고 금액", "sheet": ws.title, "formula_cell": tot, "python": float(res["total"]), "fmt": WON})
        put(ws, r, 1, "기능 유지 확인", bold=True, fill=FILL_LABEL)
        put(ws, r, 2, "동일" if sorted(sc.fp_ids) == sorted(x["id"] for x in est.fp_snapshot) else "불일치", al="center")
        put(ws, r, 3, f"시나리오 FP 항목 {len(sc.fp_ids)}건 = 기준 견적 FP 항목 {len(est.fp_snapshot)}건")
        r += 1
        for w in res["warnings"]:
            r = note(ws, r, "※ " + w, 3, F_WARN)
    page_setup(ws, False, 3)


# ---------------------------------------------------------------- FP 상세(산출내역서 'FP산정' 다단 머리글)
def _two_row_header(ws, r: int, groups: list[tuple[str, int, int]], subs: list[str], fills: dict[int, PatternFill] | None = None) -> int:
    """groups: (위 머리글, 시작열, 끝열). 한 열짜리 그룹은 위·아래를 병합. 반환: 아래 머리글 행."""
    fills = fills or {}
    for text, c1, c2 in groups:
        for cc in range(c1, c2 + 1):
            put(ws, r, cc, text if cc == c1 else "", bold=True, al="center", fill=fills.get(cc, FILL_HEAD))
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    for i, t in enumerate(subs, 1):
        put(ws, r + 1, i, t, bold=True, al="center", fill=fills.get(i, FILL_HEAD))
    for text, c1, c2 in groups:
        if c1 == c2:
            ws.cell(row=r + 1, column=c1).value = None
            ws.merge_cells(start_row=r, start_column=c1, end_row=r + 1, end_column=c1)
    ws.row_dimensions[r].height = 20
    ws.row_dimensions[r + 1].height = 30
    return r + 1


def _weight_formula(type_cell: str) -> str:
    return f'=IF({type_cell}="EI",4,IF({type_cell}="EO",5.2,IF({type_cell}="EQ",3.9,IF({type_cell}="ILF",7.5,IF({type_cell}="EIF",5.4,0)))))'


TYPES = ("ILF", "EIF", "EI", "EO", "EQ")
FILL_CALC = PatternFill("solid", fgColor="FFEB9C")  # FP 산출 열(참고 파일 ⑤가중치 열)


def _type_summary(ws, r: int, first: int, last: int, type_col: str, score_col: str, counted_col: str, py_by_type: dict, fmt: str, checks, label: str, ncols: int, weight_col: str | None = None) -> int:
    """weight_col: 정통법은 복잡도가 판정되어 가중치가 숫자인 항목만 센다(engine.calc_detailed와 같은 조건)."""
    r = section(ws, r, f"유형별 집계({label})")
    header(ws, r, ["유형", "기능 수(집계)", "FP(수식)", "설명"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=ncols)
    r += 1
    lo, hi = first, max(last, first)
    for t in TYPES:
        put(ws, r, 1, t, al="center", bold=True, fill=FILL_LABEL)
        if weight_col:
            cnt = f'=SUMPRODUCT(({type_col}{lo}:{type_col}{hi}="{t}")*({counted_col}{lo}:{counted_col}{hi}="Y")*ISNUMBER({weight_col}{lo}:{weight_col}{hi}))'
        else:
            cnt = f'=COUNTIFS({type_col}{lo}:{type_col}{hi},"{t}",{counted_col}{lo}:{counted_col}{hi},"Y")'
        put(ws, r, 2, cnt, fmt="0", formula=True, al="center")
        put(ws, r, 3, f'=SUMIFS({score_col}{lo}:{score_col}{hi},{type_col}{lo}:{type_col}{hi},"{t}")', fmt=fmt, formula=True)
        put(ws, r, 4, L.FP_TYPE_DESC[t])
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=ncols)
        bt = py_by_type[t]
        checks.append({"label": f"{label} {t} FP", "sheet": ws.title, "formula_cell": f"C{r}", "python": float(Decimal(str(bt["fp"]))), "fmt": fmt})
        checks.append({"label": f"{label} {t} 기능 수", "sheet": ws.title, "formula_cell": f"B{r}", "python": bt["count"], "fmt": "0"})
        r += 1
    return r


def _fp_simple(ws, p: Project, std, fp_sum, checks):
    items = {it.id: it for it in p.fp_items}
    ncols = 12
    widths(ws, [8, 15, 13, 24, 26, 10, 14, 6, 8, 8, 12, 30])
    sub = (f"기준: {std['id']} · 가중치 ILF 7.5 / EIF 5.4 / EI 4.0 / EO 5.2 / EQ 3.9(산정양식 수식) · 집계 Y = 확정 또는 업무 확인 · 경계 내 · 중복 아님 · "
           f"간이법은 DET·FTR 없이 유형만으로 계산 · 산정 상태: {_basis_text(fp_sum, 'simple')}")
    r = title(ws, "기능점수 산정(간이법)", sub, ncols)
    hr = _two_row_header(ws, r, [("FP ID", 1, 1), ("기능명", 2, 5), ("데이터·트랜잭션", 6, 6), ("산정 상태", 7, 8), ("FP 산출", 9, 10), ("요구사항ID", 11, 11), ("비고(분류 근거·확인 사항)", 12, 12)],
                         ["", "①애플리케이션명", "②세부 업무명", "③단위프로세스명", "단위프로세스 설명", "④FP유형", "상태", "집계", "⑤가중치", "점수", "", ""],
                         {9: FILL_CALC, 10: FILL_CALC})
    first = hr + 1
    r = hr
    for row in fp_sum["items"]:
        r += 1
        it = items[row["id"]]
        typ = row["fp_type"] or "미분류"
        put(ws, r, 1, it.id, al="center")
        put(ws, r, 2, it.application or p.overview.boundary_name or "-", al="center")
        put(ws, r, 3, it.business or "-", al="center")
        put(ws, r, 4, it.name, bold=True)
        put(ws, r, 5, it.description or "-")
        put(ws, r, 6, typ, al="center")
        put(ws, r, 7, _status_text(it, row), al="center")
        put(ws, r, 8, "Y" if row["counted"] else "N", al="center")
        put(ws, r, 9, _weight_formula(f"F{r}"), fmt=FP_FMT, formula=True)
        put(ws, r, 10, f'=IF(H{r}="Y",I{r},0)', fmt=FP_FMT, formula=True)
        put(ws, r, 11, it.requirement_ids or "-", al="center")
        put(ws, r, 12, "\n".join([f"근거: {it.rationale}" if it.rationale else "근거 미기재", *row["issues"]]))
    last = r
    r += 1
    put(ws, r, 1, "합계(집계 항목)", bold=True, al="center", fill=FILL_TOTAL)
    for cc in range(2, 10):
        put(ws, r, cc, "", fill=FILL_TOTAL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    put(ws, r, 10, f"=SUM(J{first}:J{last})" if last >= first else "=0", fmt=FP_FMT, formula=True, bold=True, fill=FILL_TOTAL)
    put(ws, r, 11, "", fill=FILL_TOTAL)
    put(ws, r, 12, "", fill=FILL_TOTAL)
    checks.append({"label": "간이법 총 FP", "sheet": ws.title, "formula_cell": f"J{r}", "python": float(Decimal(fp_sum["simple"]["total_fp"])), "fmt": FP_FMT})
    r += 1
    r = note(ws, r, "간이법 산정 상태: " + _basis_text(fp_sum, "simple") + " — " + ("; ".join(fp_sum["basis"]["simple"]["reasons"]) or "남은 항목 없음"), ncols, F_BOLD)
    r += 1
    _type_summary(ws, r, first, last, "F", "J", "H", fp_sum["simple"]["by_type"], FP_FMT, checks, "간이법", ncols)
    ws.freeze_panes = f"E{hr + 1}"  # 주소로 지정(빈 프로젝트는 이 행이 병합된 합계 행)
    if last >= first:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(ncols)}{last}"
    page_setup(ws, True, ncols, f"{hr - 1}:{hr}")


def _detailed_weight_formula(t: str, cx: str) -> str:
    return (
        f'=IF({cx}="","",IF({t}="EI",CHOOSE(MATCH({cx},{{"L","A","H"}},0),3,4,6),'
        f'IF({t}="EO",CHOOSE(MATCH({cx},{{"L","A","H"}},0),4,5,7),'
        f'IF({t}="EQ",CHOOSE(MATCH({cx},{{"L","A","H"}},0),3,4,6),'
        f'IF({t}="ILF",CHOOSE(MATCH({cx},{{"L","A","H"}},0),7,10,15),'
        f'IF({t}="EIF",CHOOSE(MATCH({cx},{{"L","A","H"}},0),5,7,10),0))))))'
    )


def _fp_detailed(ws, p: Project, std, fp_sum, checks):
    items = {it.id: it for it in p.fp_items}
    d = fp_sum["detailed"]
    ncols = 12
    widths(ws, [8, 28, 7, 7, 8, 7, 16, 14, 6, 8, 8, 40])
    sub = (f"기준: {std['id']} · 복잡도는 산정양식 VBA FPR() 표, 가중치는 FPV() 표 · DET·FTR/RET은 입력값 또는 근거 목록을 코드가 집계(*) · "
           f"산정 상태: {_basis_text(fp_sum, 'detailed')}")
    r = title(ws, "기능점수 산정(정통법)", sub, ncols)
    hr = _two_row_header(ws, r, [("FP ID", 1, 1), ("단위프로세스·데이터명", 2, 2), ("FP유형", 3, 3), ("복잡도 판정", 4, 7), ("산정 상태", 8, 9), ("FP 산출", 10, 11), ("근거·누락 정보", 12, 12)],
                         ["", "", "", "DET", "FTR/RET", "복잡도", "판정 근거", "상태", "집계", "가중치", "점수", ""],
                         {10: FILL_CALC, 11: FILL_CALC})
    first = hr + 1
    r = hr
    for row in fp_sum["items"]:
        r += 1
        it = items[row["id"]]
        derived = row["complexity_source"] == "derived"
        put(ws, r, 1, it.id, al="center")
        put(ws, r, 2, it.name, bold=True)
        put(ws, r, 3, row["fp_type"] or "미분류", al="center")
        put(ws, r, 4, row.get("det_used") if row.get("det_used") is not None else "근거 없음", al="center")
        put(ws, r, 5, row.get("ftr_ret_used") if row.get("ftr_ret_used") is not None else "근거 없음", al="center")
        put(ws, r, 6, row["complexity"] or "", al="center")
        put(ws, r, 7, L.CX_SOURCE[row["complexity_source"]] + (" *" if derived else ""), al="center")
        put(ws, r, 8, _status_text(it, row), al="center")
        put(ws, r, 9, "Y" if row["counted"] else "N", al="center")
        put(ws, r, 10, _detailed_weight_formula(f"C{r}", f"F{r}"), fmt="0", formula=True)
        put(ws, r, 11, f'=IF(AND(I{r}="Y",ISNUMBER(J{r})),J{r},0)', fmt="0", formula=True)
        put(ws, r, 12, row.get("detailed_note") or "-")
    last = r
    r += 1
    put(ws, r, 1, "합계(집계·복잡도 판정 항목)", bold=True, al="center", fill=FILL_TOTAL)
    for cc in range(2, 11):
        put(ws, r, cc, "", fill=FILL_TOTAL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    put(ws, r, 11, f"=SUM(K{first}:K{last})" if last >= first else "=0", fmt="0", formula=True, bold=True, fill=FILL_TOTAL)
    put(ws, r, 12, "", fill=FILL_TOTAL)
    checks.append({"label": "정통법 총 FP", "sheet": ws.title, "formula_cell": f"K{r}", "python": int(d["total_fp"]), "fmt": "0"})
    r += 1
    r = note(ws, r, "정통법 산정 상태: " + _basis_text(fp_sum, "detailed") + " — " + ("; ".join(fp_sum["basis"]["detailed"]["reasons"]) or "남은 항목 없음"), ncols, F_BOLD)
    if d.get("derived_ids"):
        r = note(ws, r, f"* 코드 집계 {len(d['derived_ids'])}건: 기능의 입력·출력 항목·데이터 항목 목록에서 코드가 DET·FTR/RET을 셈(FP_RULES §4-1 시스템 규칙, 원문 미확인) — 근거 목록을 검토하세요.", ncols)
    r += 1
    _type_summary(ws, r, first, last, "C", "K", "I", d["by_type"], "0", checks, "정통법", ncols, weight_col="J")
    ws.freeze_panes = f"C{hr + 1}"
    if last >= first:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(ncols)}{last}"
    page_setup(ws, True, ncols, f"{hr - 1}:{hr}")


def _factor_ref(ws, std: dict[str, Any]):
    """참조용 기준표(이 프로젝트의 산정 기준 버전 값). 계산에는 엔진의 같은 값이 쓰인다."""
    widths(ws, [16, 64, 12])
    r = title(ws, "(참조) 보정계수·가중치 기준표", f"기준 버전 {std['id']} · {std.get('title', '')} · 원문 확인 상태는 각 항목 근거 참고", 3)
    r = section(ws, r, "규모 보정계수")
    header(ws, r, ["구분", "산정방법", "근거 상태"])
    r += 1
    sf = std["size_factor"]
    put(ws, r, 1, "SW 규모", bold=True, al="center", fill=FILL_LABEL)
    put(ws, r, 2, f"= {sf['formula']} (기능점수 {sf['below_fp']} 미만 {sf['below_value']}, {sf['above_fp']} 초과 {sf['above_value']}) — {sf['basis']}")
    put(ws, r, 3, sf.get("status", ""), al="center")
    r += 2
    r = section(ws, r, "애플리케이션 복잡도 보정계수")
    header(ws, r, ["구분", "복잡도 및 난이도 수준", "보정계수"])
    r += 1
    for key, tab in std["factor_tables"].items():
        if not isinstance(tab, dict) or "levels" not in tab:
            continue
        first = r
        for lv, (txt, val) in sorted(tab["levels"].items()):
            put(ws, r, 1, tab["label"] if r == first else "", bold=True, al="center", fill=FILL_LABEL)
            put(ws, r, 2, txt)
            put(ws, r, 3, Decimal(val), fmt="0.00", al="center")
            r += 1
        if r - 1 > first:
            ws.merge_cells(start_row=first, start_column=1, end_row=r - 1, end_column=1)
    r += 1
    r = section(ws, r, "FP 유형별 가중치")
    header(ws, r, ["유형", "간이법 가중치 / 정통법 가중치(L·A·H)", "설명"])
    widths(ws, [16, 64, 12])
    r += 1
    sw = std["simple_weights"]["values"]
    dw = std["detailed_weights"]["values"]
    for t in TYPES:
        put(ws, r, 1, t, bold=True, al="center", fill=FILL_LABEL)
        put(ws, r, 2, f"간이법 {sw[t]} / 정통법 L {dw[t]['L']} · A {dw[t]['A']} · H {dw[t]['H']} — {L.FP_TYPE_DESC[t]}")
        put(ws, r, 3, "양식 수식", al="center")
        r += 1
    page_setup(ws, False, 3)


# ---------------------------------------------------------------- 설계 자료(가로 표)
def _requirements(ws, p: Project, fp_sum):
    cov = {x["requirement_id"]: x for x in fp_sum["coverage"]["rows"]}
    srcs = {s.id: s for s in p.sources}
    cols = ["요구사항ID", "구분", "분류", "요구사항명", "사용자 요구 원문", "설명", "우선\n순위", "검수 기준", "출처", "제안 주체", "상태", "FP 연결", "재검토 필요"]
    r = title(ws, "요구사항 목록", "출처 구분: 사용자 / AI 제안 / 샘플 규칙 제안 / 참고문서 추출 · 원문은 그대로 보존", len(cols))
    header(ws, r, cols, [11, 7, 10, 22, 30, 32, 6, 28, 20, 12, 8, 16, 16])
    h = r
    for q in p.requirements:
        r += 1
        src = "; ".join(f"{srcs[s].file_name} {srcs[s].location}" if s in srcs else s for s in q.source_ids) or "사용자 입력"
        fn_acc = [f"{f.id} {f.acceptance}" for f in p.functions if q.id in f.requirement_ids and f.acceptance and f.status != "excluded"]
        acc = q.acceptance or (("요구사항 기준 미정 — 기능 검수 기준: " + "; ".join(fn_acc)) if fn_acc else "검수 기준 미정")
        vals = [q.id, L.KIND[q.kind], q.category, q.title, q.original_text, q.description, L.PRIORITY[q.priority], acc, src, L.ORIGIN[q.origin], L.STATUS[q.status], cov.get(q.id, {}).get("state", ""), q.review_reason if q.needs_review else ""]
        for i, v in enumerate(vals, 1):
            put(ws, r, i, v, al="center" if i in (1, 2, 7, 11) else "top", bold=i == 4)
    finish_table(ws, h, r, len(cols))


def _functions(ws, p: Project):
    cols = ["기능ID", "기능명", "사용자", "업무 목적", "입력", "출력", "처리 규칙", "예외", "권한", "검수 기준", "요구사항ID", "화면ID", "데이터ID", "연계ID", "제안 주체", "상태"]
    r = title(ws, "기능설계", "확인되지 않은 DB·API·기존 시스템 구조는 '확인 필요'로 표시", len(cols))
    header(ws, r, cols, [8, 18, 10, 22, 20, 20, 26, 18, 12, 22, 12, 10, 10, 9, 11, 7])
    h = r
    for f in p.functions:
        r += 1
        vals = [f.id, f.name, f.actor, f.purpose, f.inputs, f.outputs, f.rules, f.exceptions, f.permissions, f.acceptance or "검수 기준 미정", f.requirement_ids or "요구사항 미연결", f.screen_ids, f.data_ids, f.interface_ids, L.ORIGIN[f.origin], L.STATUS[f.status]]
        for i, v in enumerate(vals, 1):
            put(ws, r, i, v, al="center" if i in (1, 16) else "top", bold=i == 2)
    finish_table(ws, h, r, len(cols))


def _design_lists(ws, p: Project):
    ncols = 9
    r = title(ws, "화면 · 데이터 · 연계 · 업무흐름 목록", "", ncols)
    widths(ws, [10, 24, 16, 40, 20, 18, 16, 14, 10])
    blocks = [
        ("업무흐름", ["ID", "단계명", "수행자", "설명", "요구사항ID", "상태"], [[w.id, w.name, w.actor, w.description, w.requirement_ids, L.STATUS[w.status]] for w in sorted(p.workflows, key=lambda x: x.order)]),
        ("화면 목록", ["화면ID", "화면명", "이용자", "설명", "기능ID", "상태"], [[s.id, s.name, s.users, s.description, s.function_ids, L.STATUS[s.status]] for s in p.screens]),
        ("데이터 목록·관계", ["데이터ID", "데이터명", "유지 주체", "주요 항목·관계", "하위 묶음", "개인정보", "보존기간", "기능ID", "상태"], [[d.id, d.name, {"internal": "이 시스템", "external": "외부 시스템", "unknown": "확인 필요"}[d.maintained_by], f"{d.attributes}\n관계: {d.relations}" if d.relations else d.attributes, d.subgroups, d.personal_info, d.retention or "미정", d.function_ids, L.STATUS[d.status]] for d in p.data_entities]),
        ("연계 목록", ["연계ID", "연계명", "상대 시스템", "방식·주기", "데이터 항목", "방향", "기관 확인", "기능ID", "상태"], [[i.id, i.name, i.counterpart, f"{i.method} / {i.frequency}", i.data_items, {"inbound": "수신", "outbound": "송신", "both": "양방향", "unknown": "확인 필요"}[i.direction], {"unverified": "미확인(연계 가능성 확인 필요)", "confirmed_by_org": "기관 확인", "confirmed_by_counterpart": "상대기관 확인"}[i.verified], i.function_ids, L.STATUS[i.status]] for i in p.interfaces]),
    ]
    for name, cols, rows in blocks:
        r = section(ws, r, name)
        header(ws, r, cols)
        if not rows:
            r += 1
            put(ws, r, 1, "-", al="center")
            put(ws, r, 2, "등록된 항목 없음")
        for row in rows:
            r += 1
            for i, v in enumerate(row, 1):
                put(ws, r, i, v, al="center" if i == 1 else "top", bold=i == 2)
        r += 2
    page_setup(ws, True, ncols)


def _task(ws, p: Project):
    sections = p.task_draft.sections or build_sections(p)
    sub = "과업지시서 초안(발주자 검토용). " + ("저장된 초안" if p.task_draft.sections else "저장된 초안이 없어 현재 데이터로 자동 구성")
    r = title(ws, f"{p.name} — 과업내용", sub, 2)
    header(ws, r, ["구분", "내용"], [22, 84])
    h = r
    for s in sections:
        r += 1
        put(ws, r, 1, s.title, bold=True, fill=FILL_LABEL, al="top")
        put(ws, r, 2, s.text, al="top")
    finish_table(ws, h, r, 2, landscape=False, freeze_col=1)


def _reviews(ws, p: Project):
    if not p.review_run.ran_at:
        not_run(ws, "기존 문서 검토를 실행하지 않았습니다.")
        return
    cols = ["ID", "검토 분야", "판정", "근거 위치", "원문 값", "비교 값", "판단 이유", "수정안", "업체 확인 질문", "검증 범위"]
    r = title(ws, "검토의견", f"실행 {p.review_run.ran_at} · 대상 문서 {', '.join(p.review_run.document_ids)} · 자료 일치 검토이며 실제 시스템 성능·연계·보안 검증이 아님", len(cols))
    header(ws, r, cols, [8, 14, 10, 22, 20, 20, 32, 28, 28, 16])
    h = r
    for f in p.reviews:
        r += 1
        vals = [f.id, f.category, L.VERDICT[f.verdict], f.location, f.original_value, f.compared_value, f.reason, f.suggestion, f.question_to_vendor, "자료 일치 검토" if f.doc_check_only else "문서로 확정 불가(실제 검증 필요)"]
        for i, v in enumerate(vals, 1):
            put(ws, r, i, v, al="center" if i in (1, 3) else "top")
    if not p.reviews:
        r += 1
        put(ws, r, 1, "-")
        put(ws, r, 7, "검토 결과 지적 사항 없음")
    finish_table(ws, h, r, len(cols))


def _tech(ws, p: Project):
    t = p.tech
    if not t.ran_at:
        not_run(ws, "HW·SW 구성 추천을 실행하지 않았습니다.")
        return
    srcs = {s.id: s for s in p.official_sources}
    cols = ["ID", "분야", "구성안", "선정 이유", "가정·제약", "대안", "호환성", "라이선스", "지원 상태", "운영 난도", "공식 출처·확인일", "별도 비용"]
    r = title(ws, "HW·SW 구성 추천", f"실행 {t.ran_at} · {t.mode} · 입력 가정 기반 초안이며 성능·용량을 보장하지 않음 · 장비·유료 SW 비용은 개발비에 미포함", len(cols))
    header(ws, r, cols, [8, 11, 30, 28, 28, 22, 18, 14, 18, 9, 32, 10])
    h = r
    for rec in t.recommendations:
        r += 1
        src_txt = "\n".join(f"{srcs[s].title} {srcs[s].url} (확인일 {srcs[s].checked_at or '미확인'}, {srcs[s].support_status or '상태 미확인'})" for s in rec.official_source_ids if s in srcs) or "공식 출처 미등록 — 현재 상태 미확인"
        vals = [rec.id, rec.area, rec.proposal, rec.reason, rec.assumptions, rec.alternatives, rec.compatibility, rec.license, rec.support_status, rec.ops_difficulty, src_txt, "별도 비용(개발비 미합산)" if rec.paid_item else "-"]
        for i, v in enumerate(vals, 1):
            put(ws, r, i, v, al="center" if i in (1, 2, 10) else "top")
    finish_table(ws, h, r, len(cols))


def _trace(ws, p: Project, changes: list[dict[str, Any]], fp_sum):
    ncols = 8
    r = title(ws, "근거 · 변경이력 · 미확정 사항", "", ncols)
    widths(ws, [12, 16, 14, 14, 40, 40, 22, 10])
    r = section(ws, r, "원문 출처")
    header(ws, r, ["출처ID", "문서ID", "파일명", "위치", "원문 발췌", "추출 상태"])
    h = r
    for s in p.sources:
        r += 1
        for i, v in enumerate([s.id, s.document_id, s.file_name, s.location, s.excerpt[:500], s.extraction_status], 1):
            put(ws, r, i, v, al="top")
    if not p.sources:
        r += 1
        put(ws, r, 1, "-")
        put(ws, r, 5, "참고문서 출처 없음")
    r += 2
    r = section(ws, r, "참고문서 대비 변경 내역")
    header(ws, r, ["항목ID", "문서ID", "결정", "위치", "원문", "변경 내용·이유", "기관 고유 조건", "반영 요구사항"])
    for ref in p.reference_items:
        r += 1
        dec = {"pending": "결정 대기", "keep": "유지", "modify": "수정", "exclude": "제외"}[ref.decision]
        for i, v in enumerate([ref.id, ref.document_id, dec, ref.location, ref.text, f"{ref.modified_text} / 이유: {ref.reason}" if ref.decision == "modify" else ref.reason, ref.org_specific, ref.applied_requirement_id], 1):
            put(ws, r, i, v, al="top")
    r += 2
    r = section(ws, r, "확인 목록(업무 결정·설계 가정·외부 확인·FP 검토·미실행)")
    header(ws, r, ["구분", "ID", "내용"])
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ncols)
    pend = [(x["category"], x["ref"], x["text"]) for x in cl.checklist(p, fp_sum)]
    pend += [("FP 업무 확인 전", i, "업무 확인 전 후보 — 집계 제외") for i in fp_sum["basis"]["simple"]["pending"]["not_confirmed"]]
    pend += [("FP 미연결 요구사항", i, "기능 요구사항에 연결된 FP 없음") for i in fp_sum["coverage"]["unlinked_functional"]]
    pend += [("재검토 필요", it.id, it.review_reason) for c in ("requirements", "functions", "fp_items", "screens", "data_entities", "interfaces") for it in getattr(p, c) if it.needs_review]
    for kind, iid, text in pend or [("-", "-", "확인할 항목 없음")]:
        r += 1
        put(ws, r, 1, kind, al="top")
        put(ws, r, 2, iid, al="top")
        put(ws, r, 3, text, al="top")
        for cc in range(4, ncols + 1):
            put(ws, r, cc, "")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ncols)
    r += 2
    r = section(ws, r, "변경 이력(최근 300건)")
    header(ws, r, ["시각", "주체", "대상", "ID", "변경 전", "변경 후", "영향 항목", "재검토"])
    for c in changes[:300]:
        r += 1
        vals = [c["at"], L.ORIGIN.get(c["actor"], c["actor"]), c["collection"], c["item_id"], _short(c["before"]), _short(c["after"]), c["affected"], "필요" if c["needs_review"] else ""]
        for i, v in enumerate(vals, 1):
            put(ws, r, i, v, al="top")
    ws.freeze_panes = f"A{h + 1}"
    page_setup(ws, True, ncols)


def _short(v: Any) -> str:
    if v is None:
        return ""
    s = str(v)
    return s if len(s) <= 400 else s[:400] + "…"


def _checks(ws, checks: list[dict[str, Any]]):
    cols = ["항목", "시트", "수식 셀 값", "Python 확인값", "차이", "판정"]
    r = title(ws, "검산 — Excel 수식값과 Python 계산값 비교", "수식은 Excel에서 열 때 다시 계산됩니다. '판정'이 '불일치'이면 산출물을 사용하지 말고 확인하세요.", len(cols))
    header(ws, r, cols, [30, 18, 18, 18, 14, 10])
    h = r
    for c in checks:
        r += 1
        put(ws, r, 1, c["label"])
        put(ws, r, 2, c["sheet"], al="center")
        put(ws, r, 3, f"='{c['sheet']}'!{c['formula_cell']}", fmt=c.get("fmt"), formula=True)
        if c.get("python_cell"):
            put(ws, r, 4, f"='{c['sheet']}'!{c['python_cell']}", fmt=c.get("fmt"), formula=True)
        else:
            put(ws, r, 4, c["python"], fmt=c.get("fmt"))
        put(ws, r, 5, f"=C{r}-D{r}", fmt="0.000000", formula=True)
        put(ws, r, 6, f'=IF(ABS(E{r})<0.005,"일치","불일치")', formula=True, al="center", bold=True)
    finish_table(ws, h, r, len(cols), landscape=False, freeze_col=1)


# ---------------------------------------------------------------- 미리보기
def preview(wb: Workbook, max_rows: int = 200) -> list[dict[str, Any]]:
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows)):
            vals = []
            for c in row:
                v = c.value
                if isinstance(v, str) and c.data_type == "f":
                    v = {"formula": v}
                vals.append(v)
            while vals and vals[-1] in (None, ""):
                vals.pop()
            rows.append(vals)
        out.append({"name": ws.title, "rows": rows, "truncated": ws.max_row > max_rows, "total_rows": ws.max_row})
    return out
